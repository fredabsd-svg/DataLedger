"""Leitor de LANÇAMENTOS em planilha Excel `.xlsx` (DL-077, fatia 3, frente A).

As colunas são as MESMAS do formato próprio de lançamentos (`numero;data;historico;conta;lado;
valor`, cabeçalho na linha 1 da aba `lancamentos`). A regra de agrupamento e de número é a do
formato próprio (`montar_lancamentos`); este módulo cuida da planilha.

PROTEÇÕES: as de `excel.py`, reusadas e NÃO copiadas: `_verificar_pacote` (só `.xlsx`, sem
macro, sem `.xls`, sem CSV renomeado, limite de partes e de tamanho descompactado),
`_conferir_orcamento_das_planilhas` (A4: 500 mil células e 50 por linha, contadas no stream SAX
ANTES do openpyxl, com limite por parte) e `_abrir` (duas visões: fórmulas e valores salvos).
Limite de linhas: `MAXIMO_DE_LINHAS`.

REGRAS DE CÉLULA (erro com a linha da planilha e a coluna):
- célula com erro de fórmula (`#DIV/0!`) ou fórmula sem valor salvo: erro;
- `numero`: número inteiro positivo (Excel) ou texto de 1 a 18 dígitos;
- `data`: data do Excel (sem hora) ou texto `aaaa-mm-dd` válido;
- `historico`: texto não vazio. Número é erro;
- `conta`: TEXTO. Número ou data é erro, porque o Excel converte `1.1` em `1,1`: a mensagem
  pede que a coluna seja formatada como Texto (mesma regra do plano, `excel.py`);
- `lado`: `D` ou `C`;
- `valor`: número com no máximo DUAS casas, maior que zero; ou texto `1250.00`. Mais casas
  é erro: não há arredondamento em silêncio.

Linhas totalmente vazias são ignoradas. Abas que não são `lancamentos` são ignoradas com aviso.
"""

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.contabilidade.intercambio.canonico import (
    NIVEL_AVISO,
    NIVEL_ERRO,
    Ocorrencia,
    ResultadoLeitura,
)

# Privados da fatia 1, importados e NÃO copiados: o orçamento de células (A4), a checagem da aba
# em `xl/worksheets/` (`_caminho_da_aba`, generalizada para receber o nome da aba) e o modelo
# (constantes de layout). Ver `ler` e `gerar_modelo`.
from apps.contabilidade.intercambio.formatos.excel import (
    ABA_DE_INSTRUCOES,
    LINHAS_FORMATADAS_NO_MODELO,
    _abrir,
    _caminho_da_aba,
    _conferir_orcamento_das_planilhas,
    _texto,
    _verificar_pacote,
)
from apps.contabilidade.intercambio.formatos.proprio_lancamentos_leitura import (
    CABECALHO_LANCAMENTOS,
    LinhaLida,
    _data_iso,
    montar_lancamentos,
)

FORMATO = "excel"
NOME_DA_ABA = "lancamentos"
_CENTAVO = Decimal("0.01")
_TIPO_DE_CELULA_DE_FORMULA = "f"
_TIPO_DE_CELULA_DE_ERRO = "e"


def _converter_numero(valor):
    """(texto do número, mensagem de erro ou None)."""
    if isinstance(valor, bool) or valor is None:
        return None, "número do lançamento obrigatório."
    if isinstance(valor, int):
        return (
            (str(valor), None) if valor > 0 else (None, "número do lançamento deve ser positivo.")
        )
    if isinstance(valor, float):
        if valor.is_integer() and valor > 0:
            return str(int(valor)), None
        return None, f"número '{valor}' não é inteiro positivo."
    texto = _texto(valor)
    if texto.isdigit() and len(texto) <= 18 and int(texto) > 0:
        return texto, None
    return None, f"número '{texto}': use um inteiro positivo de até 18 dígitos."


def _converter_data(valor):
    """(data, mensagem de erro ou None). Data do Excel ou texto aaaa-mm-dd."""
    if isinstance(valor, datetime):
        if valor.time().replace(microsecond=0) != datetime.min.time():
            return None, f"data com hora ({valor}): use só a data."
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    data = _data_iso(_texto(valor))
    if data is None:
        return None, f"data '{_texto(valor)}' não é a data do Excel nem aaaa-mm-dd válida."
    return data, None


def _converter_valor(valor):
    """(Decimal, mensagem de erro ou None). Número com até 2 casas, ou texto `1250.00`."""
    if isinstance(valor, bool) or valor is None:
        return None, "valor obrigatório."
    if isinstance(valor, (int, float, Decimal)):
        try:
            decimal = Decimal(str(valor))
        except ArithmeticError:
            return None, f"valor '{valor}' inválido."
        if decimal <= 0:
            return None, "valor deve ser maior que zero."
        if decimal != decimal.quantize(_CENTAVO):
            return None, f"valor {decimal} tem mais de 2 casas: não há arredondamento."
        return decimal, None
    texto = _texto(valor)
    if texto and texto.replace(".", "", 1).isdigit() and texto.count(".") == 1:
        inteiro, centavos = texto.split(".")
        if len(centavos) == 2 and Decimal(texto) > 0:
            return Decimal(texto), None
    return (
        None,
        f"valor '{texto}': use número com até 2 casas ou texto com ponto e duas casas (1250.00).",
    )


def _linha(numero, celulas, valores, ocorrencias):
    """`LinhaLida` de uma linha da planilha, ou None se ela está vazia."""
    colunas = list(CABECALHO_LANCAMENTOS)
    tamanho = max(len(celulas), len(valores), len(colunas))
    valida = True
    lidos = {}

    def recusar(coluna, mensagem):
        nonlocal valida
        ocorrencias.append(Ocorrencia(numero, coluna, NIVEL_ERRO, mensagem))
        valida = False

    alem = []
    for indice in range(tamanho):
        celula = celulas[indice] if indice < len(celulas) else None
        valor = valores[indice] if indice < len(valores) else None
        if indice >= len(colunas):
            if valor not in (None, ""):
                alem.append(indice + 1)
            continue
        coluna = colunas[indice]
        tipo = getattr(celula, "data_type", "n") if celula is not None else "n"
        if tipo == _TIPO_DE_CELULA_DE_ERRO:
            recusar(coluna, "célula com erro de fórmula (ex.: #DIV/0!).")
            lidos[coluna] = None
            continue
        if tipo == _TIPO_DE_CELULA_DE_FORMULA and valor is None:
            recusar(
                coluna,
                "fórmula sem valor calculado. Sem o valor salvo o importador não calcula: "
                "cole como valores e salve de novo.",
            )
            lidos[coluna] = None
            continue
        lidos[coluna] = valor

    if alem:
        recusar(
            f"coluna {alem[0]}",
            "a coluna está além das seis do modelo e tem conteúdo. Remova-a.",
        )
    if all(lidos.get(coluna) in (None, "") for coluna in colunas) and not alem:
        return None

    numero_texto, erro = _converter_numero(lidos.get("numero"))
    if erro:
        recusar("numero", erro)
    data, erro = _converter_data(lidos.get("data"))
    if erro:
        recusar("data", erro)

    historico = lidos.get("historico")
    if not isinstance(historico, str) or not historico.strip():
        recusar("historico", "histórico obrigatório, em texto.")
        historico = None

    conta = lidos.get("conta")
    if conta is None or isinstance(conta, bool):
        recusar("conta", "conta obrigatória.")
        conta = None
    elif not isinstance(conta, str):
        natureza = "data" if isinstance(conta, (datetime, date)) else "número"
        recusar(
            "conta",
            f"o Excel gravou o código como {natureza} ({conta}). Formate a coluna 'conta' como "
            "Texto e digite de novo: o Excel transforma 1.1 em 1,1, e o importador não converte.",
        )
        conta = None
    elif not conta.strip():
        recusar("conta", "conta obrigatória.")
        conta = None
    else:
        conta = conta.strip()

    lado = lidos.get("lado")
    lado_texto = _texto(lado) if lado is not None else ""
    if lado_texto not in ("D", "C"):
        recusar("lado", f"lado '{lado_texto}': use D (débito) ou C (crédito).")
        lado_texto = None

    valor, erro = _converter_valor(lidos.get("valor"))
    if erro:
        recusar("valor", erro)

    return LinhaLida(
        linha=numero,
        numero=numero_texto,
        data=data,
        historico=historico,
        conta=conta,
        lado=lado_texto,
        valor=valor,
        valida=valida,
    )


def _ler_planilha(formulas, valores, resultado):
    from apps.contabilidade.intercambio.leitura import MAXIMO_DE_LINHAS, ArquivoGrandeDemais

    ocorrencias = resultado.ocorrencias
    if NOME_DA_ABA not in formulas.sheetnames:
        ocorrencias.append(
            Ocorrencia(
                0,
                "aba",
                NIVEL_ERRO,
                f"a planilha não tem a aba '{NOME_DA_ABA}'. Baixe o modelo e preencha essa aba.",
            )
        )
        return
    for nome in formulas.sheetnames:
        if nome != NOME_DA_ABA:
            ocorrencias.append(
                Ocorrencia(
                    0,
                    "aba",
                    NIVEL_AVISO,
                    f"aba '{nome}' ignorada: só a aba '{NOME_DA_ABA}' é lida.",
                )
            )

    linhas = []
    pares = zip(
        formulas[NOME_DA_ABA].iter_rows(),
        valores[NOME_DA_ABA].iter_rows(values_only=True),
        strict=True,
    )
    for numero, (celulas, valor_salvo) in enumerate(pares, 1):
        if numero > MAXIMO_DE_LINHAS:
            raise ArquivoGrandeDemais(
                f"a planilha passa de {MAXIMO_DE_LINHAS} linhas. O limite é esse."
            )
        celulas = list(celulas)
        valor_salvo = list(valor_salvo)
        if numero == 1:
            cabecalho = tuple(_texto(valor) for valor in valor_salvo[: len(CABECALHO_LANCAMENTOS)])
            alem = [_texto(v) for v in valor_salvo[len(CABECALHO_LANCAMENTOS) :] if _texto(v)]
            if cabecalho != CABECALHO_LANCAMENTOS or alem:
                ocorrencias.append(
                    Ocorrencia(
                        1,
                        "cabecalho",
                        NIVEL_ERRO,
                        "a linha 1 deve ser exatamente "
                        f"{';'.join(CABECALHO_LANCAMENTOS)}, nesta ordem. "
                        "Use o modelo para baixar.",
                    )
                )
                return
            continue
        linha = _linha(numero, celulas, valor_salvo, ocorrencias)
        if linha is not None:
            linhas.append(linha)
    montar_lancamentos(linhas, resultado)


def ler(conteudo: bytes) -> ResultadoLeitura:
    """Lê os lançamentos de uma planilha `.xlsx` segura. Recusa de pacote levanta exceção."""
    _verificar_pacote(conteudo)
    resultado = ResultadoLeitura(formato=FORMATO, codificacao="")
    # Orçamento de células e colunas ANTES de o openpyxl abrir o pacote (A4, como no plano).
    # Sem esta varredura, uma planilha de lançamentos com milhões de `<c/>` era lida até o fim.
    _conferir_orcamento_das_planilhas(conteudo)
    formulas, valores = _abrir(conteudo)
    try:
        if NOME_DA_ABA in formulas.sheetnames:
            # A aba tem de estar em `xl/worksheets/`, onde o orçamento foi medido (ver `excel`).
            _caminho_da_aba(formulas, NOME_DA_ABA)
            # A dimensão declarada só serve para o openpyxl montar a aba. Sem zerá-la, o
            # conteúdo além dela seria descartado em silêncio (a nota A4 do plano).
            formulas[NOME_DA_ABA].reset_dimensions()
            valores[NOME_DA_ABA].reset_dimensions()
        _ler_planilha(formulas, valores, resultado)
    finally:
        formulas.close()
        valores.close()
    resultado.ocorrencias.sort(key=lambda o: (o.linha, o.campo))
    return resultado


# -----------------------------------------------------------------------------
# Modelo para baixar (DL-077, fatia 3, frente B)
# -----------------------------------------------------------------------------

# Posição das colunas do modelo, na ordem de `CABECALHO_LANCAMENTOS` (A = 1). Servem só para
# formatar e validar o modelo: a leitura exige o cabeçalho exato, nesta ordem.
_COLUNA_DA_DATA = CABECALHO_LANCAMENTOS.index("data") + 1
_COLUNA_DA_CONTA = CABECALHO_LANCAMENTOS.index("conta") + 1
_COLUNA_DO_LADO = CABECALHO_LANCAMENTOS.index("lado") + 1
_LARGURAS_DO_MODELO = (12, 14, 48, 22, 10, 16)

_INSTRUCOES_DE_LANCAMENTOS = (
    ("Modelo de lançamentos do DataLedger", ""),
    ("", ""),
    (
        "Como usar",
        "Preencha a aba 'lancamentos', a partir da linha 2. A linha 1 é o cabeçalho e não deve "
        "mudar. Só a aba 'lancamentos' é lida; as outras são ignoradas com aviso.",
    ),
    (
        "Uma linha por partida",
        "Cada lançamento ocupa uma linha por partida (débito ou crédito). As linhas com o mesmo "
        "número formam um lançamento, e a soma dos débitos precisa ser igual à dos créditos.",
    ),
    ("", ""),
    ("Coluna", "Valor aceito"),
    (
        "numero",
        "Inteiro positivo, de até 18 dígitos. Repete-se em cada partida do mesmo lançamento.",
    ),
    (
        "data",
        "Data do Excel, sem hora, ou texto aaaa-mm-dd válido. Repete-se em cada partida.",
    ),
    (
        "historico",
        "Texto, obrigatório. Partidas do mesmo lançamento com textos diferentes são unidas com "
        "' | ' e geram um aviso, que você aceita antes de efetivar.",
    ),
    (
        "conta",
        "Texto, obrigatório: o código da conta de origem. Formate a coluna como Texto ANTES de "
        "digitar: o Excel transforma 1.1 em 1,1, e o importador recusa número.",
    ),
    ("lado", "D = débito. C = crédito. Lista suspensa na coluna."),
    (
        "valor",
        "Número positivo com no máximo 2 casas, ou texto com ponto e duas casas (1250.00). "
        "Mais casas é erro: não há arredondamento.",
    ),
    ("", ""),
    ("Linhas totalmente vazias", "São ignoradas."),
    (
        "Fórmulas",
        "Só são aceitas com o valor calculado salvo no arquivo. Ao final, cole como valores.",
    ),
)


def gerar_modelo() -> bytes:
    """Modelo `.xlsx` vazio para os lançamentos: aba `lancamentos` e aba `instrucoes`.

    Mesmo padrão do modelo do plano (`excel.gerar_modelo`): a coluna `conta` vem formatada como
    Texto, para o código `1.1` não virar `1,1`, e o lado tem lista suspensa D/C. A aba sai sem
    nenhuma conta ou lançamento de exemplo.
    """
    pasta = Workbook()
    lancamentos = pasta.active
    lancamentos.title = NOME_DA_ABA
    lancamentos.append(list(CABECALHO_LANCAMENTOS))
    for linha in range(2, LINHAS_FORMATADAS_NO_MODELO + 1):
        lancamentos.cell(row=linha, column=_COLUNA_DA_CONTA).number_format = "@"
        lancamentos.cell(row=linha, column=_COLUNA_DA_DATA).number_format = "dd/mm/yyyy"
    for coluna, largura in zip("ABCDEF", _LARGURAS_DO_MODELO, strict=True):
        lancamentos.column_dimensions[coluna].width = largura

    validacao = DataValidation(type="list", formula1='"D,C"', allow_blank=True)
    lancamentos.add_data_validation(validacao)
    letra_do_lado = get_column_letter(_COLUNA_DO_LADO)
    validacao.add(f"{letra_do_lado}2:{letra_do_lado}{LINHAS_FORMATADAS_NO_MODELO}")

    instrucoes = pasta.create_sheet(ABA_DE_INSTRUCOES)
    for linha, (texto_a, texto_b) in enumerate(_INSTRUCOES_DE_LANCAMENTOS, start=1):
        instrucoes.cell(row=linha, column=1, value=texto_a)
        instrucoes.cell(row=linha, column=2, value=texto_b)
    instrucoes.column_dimensions["A"].width = 28
    instrucoes.column_dimensions["B"].width = 100

    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()
