"""Modelo `.xlsx` de lançamentos e os ajustes compartilhados das planilhas (DL-077, fatia 3).

Dados SINTÉTICOS. Duas coisas são provadas aqui:

1. `excel_lancamentos.gerar_modelo`: o modelo sai no padrão do plano (aba `lancamentos` com o
   cabeçalho do formato, coluna `conta` como Texto, lista suspensa D/C, aba `instrucoes`), e o
   próprio leitor o aceita sem erro. Se a coluna `conta` virar número, o leitor recusa com a
   mensagem que manda formatar como Texto.
2. Os ajustes que a integração pediu em `excel.py` (item 5): a mensagem do orçamento de células
   serve às duas planilhas e não cita o plano, e `_caminho_da_aba` recebe o nome da aba. O leitor
   de lançamentos usa a função generalizada, e não uma cópia própria.
"""

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from apps.contabilidade.intercambio.canonico import NIVEL_AVISO, NIVEL_ERRO, IntercambioRecusado
from apps.contabilidade.intercambio.formatos import excel, excel_lancamentos
from apps.contabilidade.intercambio.formatos.proprio import CABECALHO
from apps.contabilidade.intercambio.formatos.proprio_lancamentos_leitura import (
    CABECALHO_LANCAMENTOS,
)


def _pasta(abas):
    """Pasta `.xlsx` com as abas dadas: {nome: [linhas]}. A primeira aba é a ativa."""
    pasta = Workbook()
    primeira = True
    for nome, linhas in abas.items():
        aba = pasta.active if primeira else pasta.create_sheet()
        aba.title = nome
        for linha in linhas:
            aba.append(linha)
        primeira = False
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def _erros(resultado):
    return [(o.linha, o.campo, o.mensagem) for o in resultado.ocorrencias if o.nivel == NIVEL_ERRO]


# ---------------------------------------------------------------------------
# O modelo de lançamentos
# ---------------------------------------------------------------------------


def test_modelo_tem_as_duas_abas_e_o_cabecalho_do_formato_na_linha_1():
    planilha = load_workbook(io.BytesIO(excel_lancamentos.gerar_modelo()))

    assert planilha.sheetnames == ["lancamentos", "instrucoes"]
    cabecalho = [celula.value for celula in planilha["lancamentos"][1]][
        : len(CABECALHO_LANCAMENTOS)
    ]
    assert tuple(cabecalho) == CABECALHO_LANCAMENTOS
    assert planilha["lancamentos"]["A2"].value is None, "o modelo não traz conta nem lançamento"


def test_coluna_conta_do_modelo_e_texto_para_o_codigo_nao_virar_numero():
    planilha = load_workbook(io.BytesIO(excel_lancamentos.gerar_modelo()))
    aba = planilha["lancamentos"]

    coluna_da_conta = CABECALHO_LANCAMENTOS.index("conta") + 1
    letra = aba.cell(row=1, column=coluna_da_conta).column_letter
    assert aba[f"{letra}2"].number_format == "@"
    assert aba[f"{letra}1000"].number_format == "@"
    assert aba.cell(row=2, column=CABECALHO_LANCAMENTOS.index("data") + 1).number_format == (
        "dd/mm/yyyy"
    )


def test_coluna_lado_tem_lista_suspensa_de_d_e_c():
    planilha = load_workbook(io.BytesIO(excel_lancamentos.gerar_modelo()))
    aba = planilha["lancamentos"]

    validacoes = [v for v in aba.data_validations.dataValidation if isinstance(v, DataValidation)]
    lados = [v for v in validacoes if v.type == "list" and v.formula1 == '"D,C"']
    assert len(lados) == 1
    assert str(lados[0].sqref) == "E2:E1000"


def test_aba_de_instrucoes_explica_as_colunas_e_a_regra_do_debito_igual_ao_credito():
    planilha = load_workbook(io.BytesIO(excel_lancamentos.gerar_modelo()))
    textos = [celula.value for linha in planilha["instrucoes"].iter_rows() for celula in linha]
    juntos = " ".join(str(t) for t in textos if t)

    for coluna in CABECALHO_LANCAMENTOS:
        assert coluna in juntos
    assert "soma dos débitos precisa ser igual à dos créditos" in juntos
    assert "Formate a coluna como Texto" in juntos


def test_leitor_aceita_o_modelo_sem_nenhum_erro():
    resultado = excel_lancamentos.ler(excel_lancamentos.gerar_modelo())

    assert _erros(resultado) == []
    assert resultado.lancamentos == []
    # Só a aba `instrucoes` é ignorada, e isso é aviso, não erro.
    assert [(o.campo, o.nivel) for o in resultado.ocorrencias] == [("aba", NIVEL_AVISO)]


def test_modelo_preenchido_com_um_lancamento_equilibrado_e_lido_sem_erro():
    from datetime import date

    preenchido = _pasta(
        {
            "lancamentos": [
                list(CABECALHO_LANCAMENTOS),
                [1, date(2026, 1, 10), "Compra", "1.1.1", "D", 100.0],
                [1, date(2026, 1, 10), "Compra", "2.1", "C", 100.0],
            ],
            "instrucoes": [["x"]],
        }
    )

    resultado = excel_lancamentos.ler(preenchido)

    assert _erros(resultado) == []
    (lancamento,) = resultado.lancamentos
    assert lancamento.numero == "1"
    assert [(p.codigo_conta, p.lado) for p in lancamento.partidas] == [
        ("1.1.1", "debito"),
        ("2.1", "credito"),
    ]


def test_conta_gravada_como_numero_e_recusada_com_a_orientacao_de_formatar_como_texto():
    from datetime import date

    com_numero = _pasta(
        {
            "lancamentos": [
                list(CABECALHO_LANCAMENTOS),
                [1, date(2026, 1, 10), "Compra", 1.1, "D", 100.0],
                [1, date(2026, 1, 10), "Compra", "2.1", "C", 100.0],
            ]
        }
    )

    resultado = excel_lancamentos.ler(com_numero)

    erros = _erros(resultado)
    assert [campo for _linha, campo, _mensagem in erros] == ["conta"]
    assert "Formate a coluna 'conta' como Texto" in erros[0][2]


# ---------------------------------------------------------------------------
# Item 5 da integração: os ajustes compartilhados de `excel.py`
# ---------------------------------------------------------------------------


def _pasta_com_aba_bruta(aba, sheet_data):
    """Pasta `.xlsx` cuja aba tem o `sheetData` dado (XML bruto). Célula sem `r` é só isso.

    O openpyxl sempre grava `r`, então a contagem POR LINHA (mais de 50 células) só é alcançada
    com XML escrito à mão. A fatia 1 confere a referência de coluna antes dessa contagem (R3a).
    """
    base = _pasta({aba: [["x"]]})
    origem = zipfile.ZipFile(io.BytesIO(base))
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as destino:
        for parte in origem.infolist():
            dados = origem.read(parte.filename)
            if parte.filename == "xl/worksheets/sheet1.xml":
                dados = (
                    '<?xml version="1.0" encoding="UTF-8"?>'
                    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    f"<sheetData>{sheet_data}</sheetData></worksheet>"
                ).encode("utf-8")
            destino.writestr(parte.filename, dados)
    return saida.getvalue()


@pytest.mark.parametrize(
    "aba,ler",
    [
        ("plano", excel.ler),
        ("lancamentos", excel_lancamentos.ler),
    ],
)
def test_orcamento_de_celulas_por_linha_serve_as_duas_planilhas_e_nao_cita_o_plano(aba, ler):
    """Expectativa mantida: a contagem POR LINHA ("mais de 50 células") é a mesma nas duas
    planilhas e a mensagem não cita o plano. Só a CONSTRUÇÃO mudou: com `<c/>` sem `r`, a
    recusa sai pela contagem da linha, e não pela referência de coluna (que a fatia 1 vê antes)."""
    conteudo = _pasta_com_aba_bruta(aba, "<row>" + "<c/>" * 51 + "</row>")

    with pytest.raises(IntercambioRecusado) as excinfo:
        ler(conteudo)

    mensagem = excinfo.value.mensagem
    assert "mais de 50 células" in mensagem
    assert "linha 1" in mensagem
    assert "modelo do DataLedger" in mensagem
    assert "plano" not in mensagem.lower()
    assert "seis colunas" not in mensagem


@pytest.mark.parametrize(
    "aba,ler",
    [
        ("plano", excel.ler),
        ("lancamentos", excel_lancamentos.ler),
    ],
)
def test_coluna_alem_de_50_serve_as_duas_planilhas_e_nao_cita_o_plano(aba, ler):
    """A referência de coluna além de 50 (AY1, fatia 1) tem a mesma mensagem nas duas planilhas."""
    conteudo = _pasta({aba: [["x"] * 51]})

    with pytest.raises(IntercambioRecusado) as excinfo:
        ler(conteudo)

    mensagem = excinfo.value.mensagem
    assert "além da coluna 50" in mensagem
    assert "linha 1" in mensagem
    assert "modelo do DataLedger" in mensagem
    assert "plano" not in mensagem.lower()


def test_caminho_da_aba_recebe_o_nome_da_aba_e_recusa_a_aba_fora_de_xl_worksheets():
    class _Aba:
        _worksheet_path = "xl/outra/sheet1.xml"

    with pytest.raises(IntercambioRecusado) as excinfo:
        excel._caminho_da_aba({"contas_de_teste": _Aba()}, "contas_de_teste")

    assert "a aba 'contas_de_teste' não está em xl/worksheets/" in excinfo.value.mensagem


def test_caminho_da_aba_devolve_o_caminho_quando_a_aba_esta_no_lugar_certo():
    class _Aba:
        _worksheet_path = "xl/worksheets/sheet2.xml"

    assert excel._caminho_da_aba({"x": _Aba()}, "x") == "xl/worksheets/sheet2.xml"


def test_leitor_de_lancamentos_usa_a_mesma_checagem_do_plano_e_nao_tem_copia():
    """Uma só checagem de aba: a de `excel`. A cópia que existia em `excel_lancamentos` saiu."""
    assert excel_lancamentos._caminho_da_aba is excel._caminho_da_aba
    assert not hasattr(excel_lancamentos, "_conferir_caminho_da_aba")


def test_limites_de_celulas_e_colunas_sao_os_da_fatia_1():
    """Expectativa ATUALIZADA (era 500.000 células): a fatia 1 (R3, 6deb7b3) fixou o orçamento em
    60 mil células, e a fatia 3 ainda fixava o valor antigo. Quem manda é a fatia 1."""
    assert excel.MAXIMO_DE_COLUNAS_POR_LINHA == 50
    assert excel.MAXIMO_DE_CELULAS_LIDAS == 60_000
    assert CABECALHO == ("codigo", "nome", "codigo_pai", "analitica", "tipo", "natureza")
