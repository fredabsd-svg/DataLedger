"""Leitores de LANÇAMENTOS da importação (DL-077, fatia 3, frente A).

Dados SINTÉTICOS montados à mão segundo o leiaute de cada formato. Nenhuma linha de arquivo de
fornecedor ou de cliente entra aqui (RC-167). Os registros e campos citados são os do manual
oficial da ECD (pp. 130, 143-151, 63-64) e do sistema de referência (pp. 1224, 1449-1451).
"""

import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook

from apps.contabilidade.intercambio.canonico import LADO_CREDITO, LADO_DEBITO, IntercambioRecusado
from apps.contabilidade.intercambio.formatos import (
    ecd_lancamentos_leitura,
    excel_lancamentos,
    proprio_lancamentos_leitura,
    referencia_lancamentos_leitura,
)

CNPJ = "11122233000183"
PERIODO_0000 = "|0000|LECD|01012026|31012026|Empresa Sintetica Ltda|" + CNPJ + "|SP|||\r\n"


def _ecd(*linhas):
    return (PERIODO_0000 + "".join(linha + "\r\n" for linha in linhas)).encode("iso-8859-1")


def _erros(resultado):
    return [(o.linha, o.campo) for o in resultado.ocorrencias if o.nivel == "erro"]


def _avisos(resultado):
    return [(o.linha, o.campo) for o in resultado.ocorrencias if o.nivel == "aviso"]


# ---------------------------------------------------------------------------
# ECD: I200 e I250 (pp. 143-151)
# ---------------------------------------------------------------------------


def test_ecd_le_um_lancamento_normal_com_duas_partidas_e_historico_por_partida():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05012026|100,00|N||",
            "|I250|1.1.1||100,00|D|||Compra||",
            "|I250|2.1||100,00|C|||Fornecedor||",
        )
    )

    assert resultado.ocorrencias == []
    assert resultado.documento_declarado == CNPJ
    (lancamento,) = resultado.lancamentos
    assert lancamento.numero == "1"
    assert lancamento.data == date(2026, 1, 5)
    assert [(p.codigo_conta, p.lado, p.valor) for p in lancamento.partidas] == [
        ("1.1.1", LADO_DEBITO, Decimal("100.00")),
        ("2.1", LADO_CREDITO, Decimal("100.00")),
    ]
    assert [p.historico for p in lancamento.partidas] == ["Compra", "Fornecedor"]


def test_ecd_recusa_vl_lcto_diferente_da_soma_de_cada_lado_com_linha_e_campo():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05012026|100,00|N||",
            "|I250|1.1.1||90,00|D|||Compra||",
            "|I250|2.1||100,00|C|||Fornecedor||",
        )
    )

    assert resultado.lancamentos == []
    assert "VL_LCTO" in [o.campo for o in resultado.ocorrencias]


def test_ecd_lancamento_de_encerramento_e_avisado_e_nao_entra():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|3|07012026|10,00|E||",
            "|I250|1.1.1||10,00|D|||Zeramento||",
            "|I250|2.1||10,00|C|||Zeramento||",
        )
    )

    assert resultado.lancamentos == []
    assert (2, "IND_LCTO") in _avisos(resultado)
    assert _erros(resultado) == []


def test_ecd_extemporaneo_x_importa_com_aviso_de_motivo_e_exige_data_de_origem():
    importa = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|2|15012026|50,00|X|15012026|",
            "|I250|1.1.1||50,00|D|||Ajuste (motivo, data e lançamento de origem)||",
            "|I250|2.1||50,00|C|||Ajuste (motivo, data e lançamento de origem)||",
        )
    )
    assert [linha_.numero for linha_ in importa.lancamentos] == ["2"]
    assert (2, "IND_LCTO") in _avisos(importa)

    sem_data = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|2|15012026|50,00|X||",
            "|I250|1.1.1||50,00|D|||Ajuste||",
            "|I250|2.1||50,00|C|||Ajuste||",
        )
    )
    assert sem_data.lancamentos == []
    assert (2, "DT_LCTO_EXT") in _erros(sem_data)


def test_ecd_data_de_origem_so_existe_no_extemporaneo_e_ind_lcto_desconhecido_e_erro():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|2|15012026|50,00|N|15012026|",
            "|I250|1.1.1||50,00|D|||Ajuste||",
            "|I250|2.1||50,00|C|||Ajuste||",
            "|I200|4|15012026|50,00|Z||",
        )
    )

    assert (2, "DT_LCTO_EXT") in _erros(resultado)
    assert (5, "IND_LCTO") in _erros(resultado)


def test_ecd_campos_que_o_DataLedger_nao_guarda_geram_aviso_ignorado():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05012026|100,00|N||",
            "|I250|1.1.1|CC1|100,00|D|DOC/1||Compra|PART1|",
            "|I250|2.1||100,00|C|||Compra||",
        )
    )

    assert len(resultado.lancamentos) == 1
    assert {campo for _linha, campo in _avisos(resultado)} == {"COD_CCUS", "NUM_ARQ", "COD_PART"}
    assert _erros(resultado) == []


def test_ecd_historico_padronizado_usa_a_formula_do_manual_com_o_i075():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I075|H1|PAGAMENTO A FORNECEDORES|",
            "|I200|1|05012026|100,00|N||",
            "|I250|1.1.1||100,00|D||H1|complemento||",
            "|I250|2.1||100,00|C|||Fornecedor||",
        )
    )

    (lancamento,) = resultado.lancamentos
    # DESCR_HIST + " " + HIST (p. 149).
    assert lancamento.partidas[0].historico == "PAGAMENTO A FORNECEDORES complemento"


def test_ecd_cod_hist_pad_sem_i075_avisa_e_usa_so_o_hist():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05012026|100,00|N||",
            "|I250|1.1.1||100,00|D||H9|Só o HIST||",
            "|I250|2.1||100,00|C|||Fornecedor||",
        )
    )

    assert (2, "COD_HIST_PAD") in _avisos(resultado) or (3, "COD_HIST_PAD") in _avisos(resultado)
    assert resultado.lancamentos[0].partidas[0].historico == "Só o HIST"


def test_ecd_partida_sem_historico_e_erro():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05012026|100,00|N||",
            "|I250|1.1.1||100,00|D|||||",
            "|I250|2.1||100,00|C|||Fornecedor||",
        )
    )

    assert resultado.lancamentos == []
    assert any(campo == "HIST" for _linha, campo in _erros(resultado))


def test_ecd_numero_repetido_recusa_os_dois_lancamentos():
    partidas = [
        "|I250|1.1.1||10,00|D|||X||",
        "|I250|2.1||10,00|C|||X||",
    ]
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|7|05012026|10,00|N||",
            *partidas,
            "|I200|7|06012026|10,00|N||",
            *partidas,
        )
    )

    assert resultado.lancamentos == []
    assert any(campo == "NUM_LCTO" for _linha, campo in _erros(resultado))


def test_ecd_data_fora_do_periodo_do_0000_e_erro():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05022026|10,00|N||",
            "|I250|1.1.1||10,00|D|||X||",
            "|I250|2.1||10,00|C|||X||",
        )
    )

    assert resultado.lancamentos == []
    assert (2, "DT_LCTO") in _erros(resultado)


def test_ecd_valor_com_ponto_ou_sem_duas_casas_e_erro_nomeado():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I200|1|05012026|10.00|N||",
            "|I250|1.1.1||10,5|D|||X||",
            "|I250|2.1||10,00|C|||X||",
        )
    )

    assert resultado.lancamentos == []
    campos = {campo for _linha, campo in _erros(resultado)}
    assert {"VL_LCTO", "VL_DC"} <= campos


def test_ecd_0000_com_cnpj_alfanumerico_rc46_e_aceito_e_com_formato_ruim_e_erro():
    alfanumerico = "AB12CD34EF56" + "78"
    aceito = ecd_lancamentos_leitura.ler(
        (
            f"|0000|LECD|01012026|31012026|Empresa|{alfanumerico}|SP|||\r\n"
            "|I200|1|05012026|10,00|N||\r\n|I250|1.1.1||10,00|D|||X||\r\n|I250|2.1||10,00|C|||X||\r\n"
        ).encode("iso-8859-1")
    )
    assert aceito.documento_declarado == alfanumerico

    ruim = ecd_lancamentos_leitura.ler(
        "|0000|LECD|01012026|31012026|Empresa|1112223300018|SP|||\r\n".encode("iso-8859-1")
    )
    assert ruim.documento_declarado is None
    assert any(campo == "0000.6" for _linha, campo in _erros(ruim))


def test_ecd_registro_fora_do_leiaute_e_linha_sem_barras_sao_erros():
    resultado = ecd_lancamentos_leitura.ler(
        _ecd(
            "|I999|nada|",
            "I200|1|05012026|10,00|N||",
        )
    )

    campos = {campo for _linha, campo in _erros(resultado)}
    assert {"REG", "linha"} <= campos


def test_ecd_utf8_e_lido_com_aviso_de_codificacao():
    conteudo = (
        PERIODO_0000 + "|I200|1|05012026|10,00|N||\r\n"
        "|I250|1.1.1||10,00|D|||Pagamento à vista||\r\n"
        "|I250|2.1||10,00|C|||Pagamento à vista||\r\n"
    ).encode("utf-8")
    resultado = ecd_lancamentos_leitura.ler(conteudo)

    assert resultado.codificacao == "utf-8"
    assert any(campo == "codificacao" for _linha, campo in _avisos(resultado))
    assert resultado.lancamentos[0].partidas[0].historico == "Pagamento à vista"


# ---------------------------------------------------------------------------
# Formato próprio (proprio_lancamentos.py, fatia 2)
# ---------------------------------------------------------------------------

CABECALHO = "numero;data;historico;conta;lado;valor\r\n"


def _proprio(*linhas):
    return (CABECALHO + "".join(linha + "\r\n" for linha in linhas)).encode("utf-8")


def test_proprio_agrupa_partidas_pelo_numero_e_preserva_o_historico():
    resultado = proprio_lancamentos_leitura.ler(
        _proprio(
            "7;2026-03-10;Compra à vista;1.1.1;D;100.00",
            "7;2026-03-10;Compra à vista;2.1;C;100.00",
            '8;2026-03-12;"Pagamento; fornecedor X ";2.1;D;40.00',
            '8;2026-03-12;"Pagamento; fornecedor X ";1.1.1;C;40.00',
        )
    )

    assert resultado.ocorrencias == []
    assert [linha_.numero for linha_ in resultado.lancamentos] == ["7", "8"]
    assert resultado.lancamentos[1].partidas[0].historico == "Pagamento; fornecedor X "


def test_proprio_recusa_numero_em_blocos_separados_e_data_diferente_na_mesma_partida():
    resultado = proprio_lancamentos_leitura.ler(
        _proprio(
            "7;2026-03-10;A;1.1.1;D;1.00",
            "7;2026-03-10;A;2.1;C;1.00",
            "9;2026-03-10;B;1.1.1;D;2.00",
            "9;2026-03-11;B;2.1;C;2.00",
            "7;2026-03-10;A;1.1.1;D;3.00",
            "7;2026-03-10;A;2.1;C;3.00",
        )
    )

    assert resultado.lancamentos == []
    campos = {campo for _linha, campo in _erros(resultado)}
    assert {"numero", "data"} <= campos


def test_proprio_valor_com_mais_de_duas_casas_ou_virgula_e_erro():
    resultado = proprio_lancamentos_leitura.ler(
        _proprio(
            "7;2026-03-10;A;1.1.1;D;1.001",
            "7;2026-03-10;A;2.1;C;1,00",
        )
    )

    assert resultado.lancamentos == []
    assert {campo for _linha, campo in _erros(resultado)} == {"valor"}


def test_proprio_cabecalho_errado_e_erro_que_para_a_leitura():
    resultado = proprio_lancamentos_leitura.ler(
        "numero;data;historico;conta;valor;lado\r\n7;2026-03-10;A;1.1.1;1.00;D\r\n".encode("utf-8")
    )

    assert resultado.lancamentos == []
    assert resultado.ocorrencias[0].campo == "cabecalho"


# ---------------------------------------------------------------------------
# Sistema de referência: 0000, 6000 e 6100 (pp. 1224, 1225, 1449-1451)
# ---------------------------------------------------------------------------


def _referencia(*linhas):
    return ("".join(linha + "\r\n" for linha in linhas)).encode("iso-8859-1")


def _linha_6100(data="10/01/2026", debito="3", credito="12", valor="1234,56", historico="Aporte"):
    # 10 campos: REG, data, débito, crédito, valor, código do histórico, descrição, usuário,
    # filial e SCP (p. 1450). Os três últimos vazios.
    return f"|6100|{data}|{debito}|{credito}|{valor}||{historico}||||"


def test_referencia_le_x_como_um_lancamento_por_6100_com_valor_em_virgula():
    resultado = referencia_lancamentos_leitura.ler(
        _referencia(
            "|0000|" + CNPJ + "|",
            "|6000|X||||",
            _linha_6100(),
        )
    )

    assert resultado.documento_declarado == CNPJ
    assert resultado.ocorrencias == []
    (lancamento,) = resultado.lancamentos
    assert lancamento.numero == "3"  # a linha do 6100
    assert lancamento.data == date(2026, 1, 10)
    assert [(p.codigo_conta, p.lado, p.valor) for p in lancamento.partidas] == [
        ("3", LADO_DEBITO, Decimal("1234.56")),
        ("12", LADO_CREDITO, Decimal("1234.56")),
    ]


def test_referencia_lote_d_vira_um_lancamento_com_um_debito_e_varios_creditos():
    resultado = referencia_lancamentos_leitura.ler(
        _referencia(
            "|6000|D||||",
            _linha_6100(debito="10", credito="3", valor="200,00", historico="Aluguel"),
            _linha_6100(debito="10", credito="6", valor="100,00", historico="Luz"),
        )
    )

    (lancamento,) = resultado.lancamentos
    lados = [(p.codigo_conta, p.lado, p.valor) for p in lancamento.partidas]
    assert lados == [
        ("10", LADO_DEBITO, Decimal("300.00")),
        ("3", LADO_CREDITO, Decimal("200.00")),
        ("6", LADO_CREDITO, Decimal("100.00")),
    ]
    assert lancamento.numero == "1"  # linha do 6000


def test_referencia_lote_d_com_debitos_diferentes_ou_datas_diferentes_e_recusado():
    debitos_diferentes = referencia_lancamentos_leitura.ler(
        _referencia(
            "|6000|D||||",
            _linha_6100(debito="10", credito="3"),
            _linha_6100(debito="11", credito="6"),
        )
    )
    assert debitos_diferentes.lancamentos == []
    assert any(campo == "6100.3" for _linha, campo in _erros(debitos_diferentes))

    datas_diferentes = referencia_lancamentos_leitura.ler(
        _referencia(
            "|6000|C||||",
            _linha_6100(data="10/01/2026", debito="3", credito="12"),
            _linha_6100(data="11/01/2026", debito="6", credito="12"),
        )
    )
    assert datas_diferentes.lancamentos == []
    assert any(campo == "6100.2" for _linha, campo in _erros(datas_diferentes))


def test_referencia_valor_com_ponto_mais_ou_menos_de_duas_casas_e_erro_nomeado():
    resultado = referencia_lancamentos_leitura.ler(
        _referencia(
            "|6000|X||||",
            _linha_6100(valor="1234.56"),
            _linha_6100(valor="1234,5"),
            _linha_6100(valor="1234,567"),
        )
    )

    assert resultado.lancamentos == []
    assert {campo for _linha, campo in _erros(resultado)} == {"6100.5"}


def test_referencia_tipo_de_lote_fora_da_lista_e_erro_e_6100_sem_lote_tambem():
    resultado = referencia_lancamentos_leitura.ler(
        _referencia(
            "|6100|10/01/2026|3|12|1,00||Sem lote||||",
            "|6000|Z||||",
            _linha_6100(),
        )
    )

    campos = {campo for _linha, campo in _erros(resultado)}
    assert "REG" in campos
    assert "6000.2" in campos


def test_referencia_6110_filho_do_6100_gera_aviso_no_lancamento_pai():
    resultado = referencia_lancamentos_leitura.ler(
        _referencia(
            "|6000|X||||",
            _linha_6100(),
            "|6110|1|2|1,00|",
        )
    )

    assert len(resultado.lancamentos) == 1
    assert (2, "6110") in _avisos(resultado)


def test_referencia_conta_que_nao_e_codigo_reduzido_numerico_e_erro():
    resultado = referencia_lancamentos_leitura.ler(
        _referencia("|6000|X||||", _linha_6100(debito="1.1.1", credito="12"))
    )
    assert resultado.lancamentos == []
    assert any(campo == "6100.3" for _linha, campo in _erros(resultado))


def test_referencia_0000_com_cpf_ou_cnpj_so_numeros_e_aceito_e_outro_formato_e_erro():
    cpf = referencia_lancamentos_leitura.ler(_referencia("|0000|12345678909|"))
    assert cpf.documento_declarado == "12345678909"
    ruim = referencia_lancamentos_leitura.ler(_referencia("|0000|11.122.233/0001-83|"))
    assert ruim.documento_declarado is None
    assert any(campo == "0000.2" for _linha, campo in _erros(ruim))


# ---------------------------------------------------------------------------
# Excel (.xlsx): mesmas colunas, mesmas proteções do plano (excel.py)
# ---------------------------------------------------------------------------


def _xlsx(linhas, aba="lancamentos"):
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = aba
    for linha in linhas:
        planilha.append(linha)
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


CABECALHO_XLSX = ["numero", "data", "historico", "conta", "lado", "valor"]


def test_excel_le_lancamento_com_data_do_excel_e_texto_aaaa_mm_dd_e_valor_numerico():
    from datetime import datetime

    conteudo = _xlsx(
        [
            CABECALHO_XLSX,
            [7, datetime(2026, 3, 10), "Compra à vista", "1.1.1", "D", 100.0],
            [7, "2026-03-10", "Compra à vista", "2.1", "C", "100.00"],
        ]
    )

    resultado = excel_lancamentos.ler(conteudo)

    assert resultado.ocorrencias == []
    (lancamento,) = resultado.lancamentos
    assert lancamento.numero == "7"
    assert lancamento.data == date(2026, 3, 10)
    assert [(p.codigo_conta, p.lado, p.valor) for p in lancamento.partidas] == [
        ("1.1.1", LADO_DEBITO, Decimal("100")),
        ("2.1", LADO_CREDITO, Decimal("100.00")),
    ]


def test_excel_conta_numero_e_erro_e_valor_com_tres_casas_e_erro():
    conteudo = _xlsx(
        [
            CABECALHO_XLSX,
            [7, "2026-03-10", "A", 11, "D", "1.001"],
            [7, "2026-03-10", "A", "2.1", "C", 1.001],
        ]
    )

    resultado = excel_lancamentos.ler(conteudo)

    assert resultado.lancamentos == []
    campos = {campo for _linha, campo in _erros(resultado)}
    assert {"conta", "valor"} <= campos
    assert any("o Excel gravou o código como número" in o.mensagem for o in resultado.ocorrencias)


def test_excel_formula_sem_valor_salvo_e_erro_e_nao_calcula():
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = "lancamentos"
    planilha.append(CABECALHO_XLSX)
    planilha.append([7, "2026-03-10", "A", "1.1.1", "D", "=1+1"])
    saida = io.BytesIO()
    pasta.save(saida)

    resultado = excel_lancamentos.ler(saida.getvalue())

    assert resultado.lancamentos == []
    assert any(
        o.campo == "valor" and "sem valor calculado" in o.mensagem for o in resultado.ocorrencias
    )


def test_excel_recusa_xls_macro_e_pacote_que_nao_e_planilha():
    with pytest.raises(IntercambioRecusado, match=".xls"):
        excel_lancamentos.ler(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"x" * 200)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as pacote:
        pacote.writestr("[Content_Types].xml", "<Types/>")
        pacote.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(IntercambioRecusado, match="macros"):
        excel_lancamentos.ler(buffer.getvalue())

    with pytest.raises(IntercambioRecusado, match="planilha"):
        excel_lancamentos.ler(b"numero;data\r\n1;2026-01-01\r\n")


def test_excel_sem_aba_lancamentos_e_erro_e_aba_extra_e_aviso():
    sem_aba = excel_lancamentos.ler(_xlsx([CABECALHO_XLSX], aba="plano"))
    assert any(o.campo == "aba" and o.nivel == "erro" for o in sem_aba.ocorrencias)

    pasta = Workbook()
    pasta.active.title = "lancamentos"
    pasta.active.append(CABECALHO_XLSX)
    pasta.active.append([7, "2026-03-10", "A", "1.1.1", "D", "1.00"])
    pasta.active.append([7, "2026-03-10", "A", "2.1", "C", "1.00"])
    outra = pasta.create_sheet("notas")
    outra.append(["x"])
    saida = io.BytesIO()
    pasta.save(saida)
    com_extra = excel_lancamentos.ler(saida.getvalue())
    assert any(o.campo == "aba" and o.nivel == "aviso" for o in com_extra.ocorrencias)
    assert len(com_extra.lancamentos) == 1


def test_excel_cabecalho_errado_e_erro():
    resultado = excel_lancamentos.ler(
        _xlsx([["numero", "data", "historico", "conta", "valor", "lado"]])
    )
    assert resultado.lancamentos == []
    assert any(o.campo == "cabecalho" for o in resultado.ocorrencias)


def test_limite_de_bytes_do_leitor_e_o_mesmo_do_plano():
    from apps.contabilidade.intercambio.importacao_lancamentos import _conferir_limites
    from apps.contabilidade.intercambio.leitura import (
        TAMANHO_MAXIMO_ARQUIVO_BYTES,
        ArquivoGrandeDemais,
    )

    with pytest.raises(ArquivoGrandeDemais):
        _conferir_limites(b"x" * (TAMANHO_MAXIMO_ARQUIVO_BYTES + 1))
