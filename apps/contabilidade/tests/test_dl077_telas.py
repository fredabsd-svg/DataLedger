"""DL-077, fatia 1, frente C: telas de importação e exportação do plano de contas.

Chamadas reais pelo cliente HTTP com sessão autenticada (o mesmo caminho do
contador). Dados SINTÉTICOS: escritórios, empresas e CNPJs fictícios.

As contagens e as ocorrências esperadas foram DERIVADAS À MÃO do leiaute de cada
formato e das regras do núcleo (`apps.contabilidade.intercambio.plano`), antes de
rodar: um teste que só repete a saída do código não prova nada. Quando a saída
diverge do esperado, a divergência é investigada, não copiada para o teste.

Permissões (as da API): importar e aplicar exigem `PodeEscriturar` (ADMINISTRADOR,
GESTOR, ANALISTA, FINANCEIRO); exportar exige a leitura (também PARALEGAL); CLIENTE
não passa em nenhuma das rotas.
"""

import hashlib
import io
import re
from datetime import date

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.auditoria.models import RegistroAuditoria
from apps.contabilidade.intercambio.canonico import NIVEL_AVISO, NIVEL_ERRO
from apps.contabilidade.intercambio.formatos import ESCRITORES, LEITORES
from apps.contabilidade.intercambio.formatos.proprio import CABECALHO
from apps.contabilidade.intercambio.plano import FILTROS_DE_EXPORTACAO, POLITICAS
from apps.contabilidade.models import Conta, NaturezaConta, TipoConta
from apps.contabilidade.tests.test_dl024_atalhos_e_acessibilidade import assert_pagina_acessivel
from apps.contabilidade.views_web import (
    FILTROS_DE_EXPORTACAO_NA_TELA,
    FORMATOS_DE_EXPORTACAO,
    FORMATOS_DE_IMPORTACAO,
    POLITICAS_DE_IMPORTACAO_NA_TELA,
)
from apps.empresas.models import Empresa, ModoEscrituracao
from apps.empresas.services import MENSAGEM_RECUSA_CONTABILIDADE_LIVRO_CAIXA
from apps.tenancy.models import Escritorio, Papel, VinculoUsuarioEscritorio

pytestmark = pytest.mark.django_db

SENHA = "senha-forte-123"
CNPJ_DA_EMPRESA = "44444444000144"  # sintético
CNPJ_DE_OUTRA_EMPRESA = "55555555000155"  # sintético
DATA_FIXA = date(2026, 10, 8)

PADRAO_CAMPO_OCULTO = re.compile(r'<input type="hidden" name="([^"]+)" value="([^"]*)">')
PADRAO_SHA256 = re.compile(r"^[0-9a-f]{64}$")

# ---------------------------------------------------------------------------
# Arquivos sintéticos, montados à mão segundo o leiaute de cada formato.
# ---------------------------------------------------------------------------

# ECD (I050, ISO-8859-1, CRLF). COD_NAT_CC 01 traz o tipo (ativo); 04 não traz (HI-87).
# Natureza nunca vem na ECD (HI-88): toda conta com tipo recebe natureza presumida.
ECD = (
    "|I050|01012023|01|S|1|1||Ativo|\r\n"
    "|I050|01012023|01|A|2|1.1|1|Caixa|\r\n"
    "|I050|01012023|04|S|1|4||Despesas|\r\n"
    "|I050|01012023|04|A|2|4.1|4|Aluguel|\r\n"
).encode("iso-8859-1")

# DataLedger TXT (UTF-8, `;`, cabeçalho). Conta 1.1 herda o tipo do pai.
PROPRIO = (
    "codigo;nome;codigo_pai;analitica;tipo;natureza\r\n"
    "1;Ativo;;N;ativo;devedora\r\n"
    "1.1;Caixa;1;S;;\r\n"
    "3;Receitas;;N;;\r\n"
    "3.1;Vendas;3;S;;\r\n"
    "9;Outros;;S;;\r\n"
).encode("utf-8")

# Sistema de referência (ISO-8859-1, `|`, dd/mm/aaaa). Forma CANÔNICA: `|` no início e no
# fim de cada registro (decisão do arquiteto, 08/10/2026). Sem tipo no leiaute: todo tipo
# vem de prefixo ou de conta superior. A conta "3" nasce inativa (situação I) e o leitor
# avisa disso (0200.7).
REFERENCIA = (
    f"|0000|{CNPJ_DA_EMPRESA}|\r\n"
    "|0200|1|1|S|Ativo|01/01/2023|A|||||\r\n"
    "|0200|2|1.1|A|Caixa|01/01/2023|A|||||\r\n"
    "|0200|3|3|S|Receitas|01/01/2023|I|01/06/2024||||\r\n"
).encode("iso-8859-1")


def _planilha(linhas):
    pasta = Workbook()
    aba = pasta.active
    aba.title = "plano"
    for linha in linhas:
        aba.append(linha)
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


# Planilha no modelo: cabeçalho do formato próprio, código como TEXTO.
PLANILHA = _planilha(
    [
        list(CABECALHO),
        ["1", "Ativo", None, "N", "ativo", "devedora"],
        ["1.1", "Caixa", "1", "S", None, None],
        ["9", "Outros", None, "S", None, None],
    ]
)
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# (formato, conteúdo, nome, contagens esperadas, erros {(linha, campo)}, avisos {(linha, campo)})
CASOS_DE_PREVIA = [
    pytest.param(
        "ecd",
        ECD,
        "plano.txt",
        {"criar": 2, "atualizar": 0, "sem_mudanca": 0, "recusada": 2},
        {(3, "tipo"), (4, "codigo_pai"), (4, "tipo")},
        # A12: duas contas novas sem classificação (linha 0). A9: o trecho não traz 0000.
        {(0, "0000"), (0, "classificacao"), (1, "natureza"), (2, "natureza")},
        id="ecd",
    ),
    pytest.param(
        "proprio",
        PROPRIO,
        "plano.txt",
        {"criar": 2, "atualizar": 0, "sem_mudanca": 0, "recusada": 3},
        {(4, "tipo"), (5, "codigo_pai"), (5, "tipo"), (6, "tipo")},
        # A12: duas contas novas sem classificação (linha 0).
        {(0, "classificacao"), (3, "natureza")},
        id="proprio",
    ),
    pytest.param(
        "referencia",
        REFERENCIA,
        "plano.txt",
        {"criar": 0, "atualizar": 0, "sem_mudanca": 0, "recusada": 3},
        {(2, "tipo"), (3, "codigo_pai"), (3, "tipo"), (4, "tipo")},
        # Conta inativa no arquivo gera aviso no leiaute de referência (campo 0200.7).
        {(4, "0200.7")},
        id="referencia",
    ),
    pytest.param(
        "excel",
        PLANILHA,
        "plano.xlsx",
        {"criar": 2, "atualizar": 0, "sem_mudanca": 0, "recusada": 1},
        {(4, "tipo")},
        # A12: duas contas novas sem classificação (linha 0).
        {(0, "classificacao"), (3, "natureza")},
        id="excel",
    ),
]

# (formato, conteúdo, nome, prefixos da tela, contas criadas depois de aplicar)
CASOS_DE_APLICACAO = [
    pytest.param(
        "ecd",
        ECD,
        "plano.txt",
        {"4": "despesa"},
        {"1", "1.1", "4", "4.1"},
        id="ecd",
    ),
    pytest.param(
        "proprio",
        PROPRIO,
        "plano.txt",
        {"3": "receita", "9": "despesa"},
        {"1", "1.1", "3", "3.1", "9"},
        id="proprio",
    ),
    pytest.param(
        "referencia",
        REFERENCIA,
        "plano.txt",
        {"1": "ativo", "3": "receita"},
        {"1", "1.1", "3"},
        id="referencia",
    ),
    pytest.param(
        "excel",
        PLANILHA,
        "plano.xlsx",
        {"9": "despesa"},
        {"1", "1.1", "9"},
        id="excel",
    ),
]


# ---------------------------------------------------------------------------
# Cenário
# ---------------------------------------------------------------------------


def _usuario(username, escritorio, papel):
    usuario = get_user_model().objects.create_user(
        username=username, email=f"{username}@escritorio.com.br", password=SENHA
    )
    VinculoUsuarioEscritorio.objects.create(usuario=usuario, escritorio=escritorio, papel=papel)
    return usuario


@pytest.fixture
def cenario():
    escritorio_a = Escritorio.objects.create(nome="Escritório Plano A", cnpj="11111111000111")
    escritorio_b = Escritorio.objects.create(nome="Escritório Plano B", cnpj="22222222000122")
    empresa = Empresa.objects.create(
        escritorio=escritorio_a,
        razao_social="Empresa Plano Ltda",
        cnpj=CNPJ_DA_EMPRESA,
    )
    outra_da_mesma = Empresa.objects.create(
        escritorio=escritorio_a,
        razao_social="Outra Plano Ltda",
        cnpj="66666666000166",
    )
    outra_de_outro_escritorio = Empresa.objects.create(
        escritorio=escritorio_b,
        razao_social="Empresa de Outro Escritório Ltda",
        cnpj=CNPJ_DE_OUTRA_EMPRESA,
    )
    _usuario("gestor-plano", escritorio_a, Papel.GESTOR)
    _usuario("analista-plano", escritorio_a, Papel.ANALISTA)
    _usuario("financeiro-plano", escritorio_a, Papel.FINANCEIRO)
    _usuario("paralegal-plano", escritorio_a, Papel.PARALEGAL)
    _usuario("cliente-plano", escritorio_a, Papel.CLIENTE)
    _usuario("gestor-outro-escritorio", escritorio_b, Papel.GESTOR)
    return {
        "empresa": empresa,
        "outra_da_mesma": outra_da_mesma,
        "de_outro_escritorio": outra_de_outro_escritorio,
    }


@pytest.fixture
def empresa_livro_caixa(cenario):
    """Mesmo escritório do gestor-plano, empresa em modo livro-caixa (DL-038)."""
    escritorio = cenario["empresa"].escritorio
    return Empresa.objects.create(
        escritorio=escritorio,
        razao_social="Empresa Livro-Caixa Ltda",
        cnpj="77777777000177",
        modo_escrituracao=ModoEscrituracao.LIVRO_CAIXA,
    )


@pytest.fixture
def relogio(monkeypatch):
    """Data fixa para o nome do arquivo exportado (o nome leva a data de hoje)."""
    monkeypatch.setattr(timezone, "localdate", lambda *a, **k: DATA_FIXA)


def _entrar(client, username):
    assert client.login(username=username, password=SENHA)


def _url(nome, empresa, **kwargs):
    return reverse(f"contabilidade_web:{nome}", args=[empresa.id], kwargs=kwargs or None)


def _arquivo(conteudo, nome):
    return SimpleUploadedFile(nome, conteudo, content_type="application/octet-stream")


def _previa(client, empresa, conteudo, nome, formato, politica="so_acrescentar", **extra):
    return client.post(
        _url("plano_importar", empresa),
        {"arquivo": _arquivo(conteudo, nome), "formato": formato, "politica": politica, **extra},
    )


def _campos_ocultos(html):
    return dict(PADRAO_CAMPO_OCULTO.findall(html))


def _aplicar(client, empresa, html_da_previa, conteudo, nome):
    """Aplica como o botão da tela: os campos ocultos da prévia + o arquivo reenviado."""
    campos = _campos_ocultos(html_da_previa)
    return client.post(
        _url("plano_importar_aplicar", empresa), {**campos, "arquivo": _arquivo(conteudo, nome)}
    )


def _contas(empresa):
    return set(Conta.objects.filter(empresa=empresa).values_list("codigo", flat=True))


def _erros(previa):
    return {(o.linha, o.campo) for o in previa.ocorrencias if o.nivel == NIVEL_ERRO}


def _avisos(previa):
    return {(o.linha, o.campo) for o in previa.ocorrencias if o.nivel == NIVEL_AVISO}


# ---------------------------------------------------------------------------
# Formulário e prévia
# ---------------------------------------------------------------------------


def test_formulario_de_importacao_mostra_formatos_politica_e_modelo(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(_url("plano_importar", cenario["empresa"]))

    assert resposta.status_code == 200
    html = resposta.content.decode()
    for rotulo in FORMATOS_DE_IMPORTACAO.values():
        assert rotulo in html
    assert POLITICAS_DE_IMPORTACAO_NA_TELA["so_acrescentar"] in html
    assert "Baixar o modelo da planilha Excel" in html
    assert _url("plano_modelo_excel", cenario["empresa"]) in html
    assert "A importação nunca apaga conta" in html


@pytest.mark.parametrize("formato,conteudo,nome,contagens,erros,avisos", CASOS_DE_PREVIA)
def test_previa_de_cada_formato_tem_contagens_e_ocorrencias_escritas_a_mao(
    client, cenario, formato, conteudo, nome, contagens, erros, avisos
):
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], conteudo, nome, formato)

    assert resposta.status_code == 200, resposta.content.decode()
    previa = resposta.context["previa"]
    assert previa.contagens == contagens
    assert _erros(previa) == erros
    assert _avisos(previa) == avisos
    assert previa.sha256 == hashlib.sha256(conteudo).hexdigest()
    # A prévia não grava: nada no cadastro.
    assert _contas(cenario["empresa"]) == set()


def test_previa_com_erro_nao_mostra_o_botao_aplicar(client, cenario):
    """Mutante-alvo do critério 3: se o botão Aplicar aparecer com erro, este teste cai."""
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], PROPRIO, "plano.txt", "proprio")

    html = resposta.content.decode()
    assert "Não pode aplicar: há erro" in html
    assert "Aplicar plano de contas" not in html
    assert _url("plano_importar_aplicar", cenario["empresa"]) not in html
    assert 'name="assinatura"' not in html
    assert "Aplicar fica indisponível enquanto houver erro" in html


def test_previa_sem_erro_mostra_o_botao_aplicar_com_sha_e_assinatura(client, cenario):
    _entrar(client, "gestor-plano")

    # Prefixo "4" resolve as contas sem tipo; sem ele o caso é de erro (tabela acima).
    resposta = _previa(
        client, cenario["empresa"], ECD, "plano.txt", "ecd", prefixo_1="4", tipo_1="despesa"
    )

    html = resposta.content.decode()
    previa = resposta.context["previa"]
    assert not previa.tem_erro
    assert "Pode aplicar: não há erro" in html
    assert "Aplicar plano de contas" in html
    campos = _campos_ocultos(html)
    assert campos["sha256"] == previa.sha256
    assert PADRAO_SHA256.match(campos["sha256"])
    assert campos["assinatura"] == previa.assinatura
    assert PADRAO_SHA256.match(campos["assinatura"])


def test_avisos_de_natureza_presumida_listam_cada_conta_para_conferir(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], PROPRIO, "plano.txt", "proprio")

    html = resposta.content.decode()
    assert "Conferir a natureza." in html
    assert "1.1 — Caixa" in html  # presumida pelo tipo herdado (ativo -> devedora)
    assert "depreciação acumulada" in html


def test_contas_inativas_do_arquivo_aparecem_destacadas(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], REFERENCIA, "plano.txt", "referencia")

    html = resposta.content.decode()
    assert "Contas inativas no arquivo." in html
    assert "3 — Receitas" in html
    assert "O código reduzido do arquivo não é gravado" in html


def test_arquivo_sem_conta_vira_ocorrencia_do_arquivo_inteiro(client, cenario):
    """Linha 0 é o arquivo inteiro: arquivo sem nenhuma conta vira ocorrência do arquivo."""
    _entrar(client, "gestor-plano")
    vazio = ("codigo;nome;codigo_pai;analitica;tipo;natureza\r\n").encode("utf-8")

    resposta = _previa(client, cenario["empresa"], vazio, "plano.txt", "proprio")

    previa = resposta.context["previa"]
    assert {(o.linha, o.campo) for o in previa.ocorrencias} == {(0, "arquivo")}
    assert "o arquivo não traz nenhuma conta do plano." in resposta.content.decode()


# ---------------------------------------------------------------------------
# Prefixos de tipo
# ---------------------------------------------------------------------------


def test_prefixos_mudam_o_tipo_e_a_previa_passa_a_poder_aplicar(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(
        client,
        cenario["empresa"],
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="9",
        tipo_2="despesa",
    )

    previa = resposta.context["previa"]
    assert not previa.tem_erro
    assert previa.contagens == {"criar": 5, "atualizar": 0, "sem_mudanca": 0, "recusada": 0}
    por_codigo = {i.codigo: i for i in previa.itens}
    assert (por_codigo["3"].tipo, por_codigo["3"].origem_tipo) == ("receita", "prefixo")
    # 3.1 não tem tipo no arquivo: herda o da conta superior, que veio do prefixo.
    assert (por_codigo["3.1"].tipo, por_codigo["3.1"].origem_tipo) == (
        "receita",
        "conta_superior",
    )
    assert (por_codigo["9"].tipo, por_codigo["9"].natureza) == ("despesa", "devedora")
    html = resposta.content.decode()
    assert "Pode aplicar: não há erro" in html
    # Os prefixos usados aparecem na tela, com o rótulo do tipo, para o contador conferir.
    assert "3 → Receita" in html
    assert "9 → Despesa" in html


def test_prefixo_repetido_e_tipo_inexistente_sao_recusados_na_tela(client, cenario):
    _entrar(client, "gestor-plano")

    repetido = _previa(
        client,
        cenario["empresa"],
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="3",
        tipo_2="despesa",
    )
    tipo_invalido = _previa(
        client,
        cenario["empresa"],
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="lucro",
    )
    prefixo_sem_tipo = _previa(
        client, cenario["empresa"], PROPRIO, "plano.txt", "proprio", prefixo_1="3"
    )

    assert repetido.status_code == 400
    assert "aparece mais de uma vez" in repetido.content.decode()
    assert tipo_invalido.status_code == 400
    assert "não existe" in tipo_invalido.content.decode()
    assert prefixo_sem_tipo.status_code == 400
    assert "informe o começo do código e o tipo" in prefixo_sem_tipo.content.decode()
    assert _contas(cenario["empresa"]) == set()


def test_formulario_de_prefixo_so_aparece_quando_ha_conta_sem_tipo(client, cenario):
    _entrar(client, "gestor-plano")

    com_tipo_faltando = _previa(client, cenario["empresa"], PROPRIO, "plano.txt", "proprio")
    sem_tipo_faltando = _previa(
        client,
        cenario["empresa"],
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="9",
        tipo_2="despesa",
    )

    assert "Tipo por prefixo" in com_tipo_faltando.content.decode()
    assert "Tipo por prefixo" not in sem_tipo_faltando.content.decode()


# ---------------------------------------------------------------------------
# Aplicação
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("formato,conteudo,nome,prefixos,esperadas", CASOS_DE_APLICACAO)
def test_aplicar_cria_as_contas_de_cada_formato_depois_do_prefixo(
    client, cenario, formato, conteudo, nome, prefixos, esperadas
):
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    # Sem prefixo, o arquivo tem erro: a conferência é a da tabela de prévias.
    sem_prefixo = _previa(client, empresa, conteudo, nome, formato)
    assert sem_prefixo.context["previa"].tem_erro

    extras = {}
    for indice, (prefixo, tipo) in enumerate(prefixos.items(), start=1):
        extras[f"prefixo_{indice}"] = prefixo
        extras[f"tipo_{indice}"] = tipo
    com_prefixo = _previa(client, empresa, conteudo, nome, formato, **extras)
    assert not com_prefixo.context["previa"].tem_erro, com_prefixo.context["previa"].ocorrencias

    resposta = _aplicar(client, empresa, com_prefixo.content.decode(), conteudo, nome)

    assert resposta.status_code == 302, resposta.content.decode()
    assert resposta.url == _url("plano_de_contas", empresa)
    assert _contas(empresa) == esperadas


def test_aplicar_mostra_a_contagem_na_tela_do_plano(client, cenario):
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    previa = _previa(
        client,
        empresa,
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="9",
        tipo_2="despesa",
    )

    resposta = _aplicar(client, empresa, previa.content.decode(), PROPRIO, "plano.txt")
    seguida = client.get(resposta.url)

    html = seguida.content.decode()
    assert "Contas criadas: 5." in html
    assert "Nomes atualizados: 0." in html
    assert "Já existiam sem mudança: 0." in html


def test_aplicar_grava_trilha_de_quem_importou_com_sha_e_contagens(client, cenario):
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    previa = _previa(
        client,
        empresa,
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="9",
        tipo_2="despesa",
    )

    _aplicar(client, empresa, previa.content.decode(), PROPRIO, "plano.txt")

    registro = RegistroAuditoria.objects.get(acao="plano_de_contas.importado")
    assert registro.detalhes["sha256"] == hashlib.sha256(PROPRIO).hexdigest()
    assert registro.detalhes["criadas"] == 5
    assert registro.detalhes["formato"] == "proprio"


def test_aplicar_com_erro_direto_e_recusado_sem_gravar_nada(client, cenario):
    """Formulário montado à mão com o SHA-256 de um arquivo com erro: recusa, nada gravado."""
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    previa = _previa(client, empresa, PROPRIO, "plano.txt", "proprio")
    campos = _campos_ocultos(previa.content.decode())
    assert "sha256" not in campos  # a prévia com erro não entrega o formulário de aplicar

    resposta = client.post(
        _url("plano_importar_aplicar", empresa),
        {
            "formato": "proprio",
            "politica": "so_acrescentar",
            "sha256": previa.context["previa"].sha256,
            "assinatura": previa.context["previa"].assinatura,
            "arquivo": _arquivo(PROPRIO, "plano.txt"),
        },
    )

    assert resposta.status_code == 400
    assert "Nada foi gravado" in resposta.content.decode()
    assert _contas(empresa) == set()
    assert not RegistroAuditoria.objects.filter(acao="plano_de_contas.importado").exists()


def test_aplicar_com_arquivo_diferente_do_conferido_e_recusado_sem_gravar(client, cenario):
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    previa_a = _previa(client, empresa, ECD, "plano.txt", "ecd", prefixo_1="4", tipo_1="despesa")
    outro_arquivo = (
        "|I050|01012023|01|S|1|1||Ativo|\r\n|I050|01012023|01|A|2|1.1|1|Bancos|\r\n"
    ).encode("iso-8859-1")

    resposta = _aplicar(client, empresa, previa_a.content.decode(), outro_arquivo, "plano.txt")

    assert resposta.status_code == 409
    html = resposta.content.decode()
    assert "não é o que foi conferido na prévia" in html
    assert "Nada foi gravado" in html
    assert _contas(empresa) == set()


def test_aplicar_depois_de_o_cadastro_mudar_e_recusado_sem_gravar(client, cenario):
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    previa = _previa(
        client,
        empresa,
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="9",
        tipo_2="despesa",
    )
    # Entre a prévia e a aplicação, a conta "1" aparece no cadastro: a conferência muda.
    Conta.objects.create(
        empresa=empresa,
        codigo="1",
        nome="Ativo",
        tipo=TipoConta.ATIVO,
        natureza=NaturezaConta.DEVEDORA,
        aceita_lancamento=False,
    )

    resposta = _aplicar(client, empresa, previa.content.decode(), PROPRIO, "plano.txt")

    assert resposta.status_code == 409
    assert "o cadastro mudou desde a prévia" in resposta.content.decode()
    assert _contas(empresa) == {"1"}


def test_aplicar_sem_o_resumo_da_previa_e_recusado(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.post(
        _url("plano_importar_aplicar", cenario["empresa"]),
        {"formato": "proprio", "politica": "so_acrescentar", "arquivo": _arquivo(PROPRIO, "p.txt")},
    )

    assert resposta.status_code == 400
    assert "Faça a conferência de novo" in resposta.content.decode()
    assert _contas(cenario["empresa"]) == set()


def test_conferencia_ignora_contas_de_outra_empresa_do_mesmo_escritorio(client, cenario):
    """Isolamento: a conta "1" de outra empresa do escritório não é "existente" aqui."""
    _entrar(client, "gestor-plano")
    Conta.objects.create(
        empresa=cenario["outra_da_mesma"],
        codigo="1",
        nome="Passivo de outra",
        tipo=TipoConta.PASSIVO,
        natureza=NaturezaConta.CREDORA,
        aceita_lancamento=False,
    )

    resposta = _previa(
        client,
        cenario["empresa"],
        PROPRIO,
        "plano.txt",
        "proprio",
        prefixo_1="3",
        tipo_1="receita",
        prefixo_2="9",
        tipo_2="despesa",
    )

    previa = resposta.context["previa"]
    por_codigo = {i.codigo: i for i in previa.itens}
    assert por_codigo["1"].acao == "criar"
    assert por_codigo["1"].tipo == "ativo"


# ---------------------------------------------------------------------------
# Limites da entrada
# ---------------------------------------------------------------------------


def test_formato_desconhecido_e_recusado(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], PROPRIO, "plano.txt", "xpto")

    assert resposta.status_code == 400
    assert "Escolha o formato do arquivo." in resposta.content.decode()


def test_planilha_com_macro_xlsm_e_recusada_com_mensagem(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], PLANILHA, "plano.xlsm", "excel")

    assert resposta.status_code == 400
    assert "não é aceito" in resposta.content.decode()


def test_campo_nao_contratado_na_previa_e_recusado(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], PROPRIO, "plano.txt", "proprio", xpto="x")

    assert resposta.status_code == 400
    assert _contas(cenario["empresa"]) == set()


def test_planilha_corrompida_e_recusada_com_mensagem_e_nao_da_500(client, cenario):
    """O núcleo recusa o conteúdo que não é .xlsx (IntercambioRecusado): a tela mostra a
    recusa com 400, nunca uma exceção."""
    _entrar(client, "gestor-plano")

    resposta = _previa(client, cenario["empresa"], b"isto nao e um xlsx", "plano.xlsx", "excel")

    assert resposta.status_code == 400
    assert "não é uma planilha .xlsx" in resposta.content.decode()
    assert _contas(cenario["empresa"]) == set()


def test_arquivo_ausente_e_recusado(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.post(
        _url("plano_importar", cenario["empresa"]),
        {"formato": "proprio", "politica": "so_acrescentar"},
    )

    assert resposta.status_code == 400
    assert "Escolha o arquivo do plano de contas." in resposta.content.decode()


def test_politica_que_nao_existe_e_recusada(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = _previa(
        client, cenario["empresa"], PROPRIO, "plano.txt", "proprio", politica="apaga-tudo"
    )

    assert resposta.status_code == 400
    assert "Escolha a política de importação." in resposta.content.decode()


# ---------------------------------------------------------------------------
# Exportação
# ---------------------------------------------------------------------------


def _plano_de_duas_contas(empresa):
    Conta.objects.create(
        empresa=empresa,
        codigo="1",
        nome="Ativo",
        tipo=TipoConta.ATIVO,
        natureza=NaturezaConta.DEVEDORA,
        aceita_lancamento=False,
    )
    Conta.objects.create(
        empresa=empresa,
        codigo="1.1",
        nome="Caixa",
        tipo=TipoConta.ATIVO,
        natureza=NaturezaConta.DEVEDORA,
        aceita_lancamento=True,
        conta_pai=Conta.objects.get(empresa=empresa, codigo="1"),
    )


def test_formulario_de_exportacao_mostra_os_avisos_fixos(client, cenario):
    _entrar(client, "paralegal-plano")

    resposta = client.get(_url("plano_exportar", cenario["empresa"]))

    assert resposta.status_code == 200
    html = resposta.content.decode()
    assert (
        "Arquivo com os registros do plano no leiaute da ECD — não é a ECD: não tem os "
        "demais blocos, termos nem assinatura e não passou pelo programa da Receita."
    ) in html
    assert "Código reduzido sequencial pela ordem do código: não é estável" in html
    for rotulo in FORMATOS_DE_EXPORTACAO.values():
        assert rotulo[0] in html
    for rotulo in FILTROS_DE_EXPORTACAO_NA_TELA.values():
        assert rotulo in html


def test_exportacao_no_formato_proprio_baixa_o_arquivo_com_nome_e_conteudo(
    client, cenario, relogio
):
    _entrar(client, "gestor-plano")
    _plano_de_duas_contas(cenario["empresa"])

    resposta = client.get(
        _url("plano_exportar", cenario["empresa"]), {"formato": "proprio", "filtro": "todas"}
    )

    assert resposta.status_code == 200
    assert resposta["Content-Disposition"] == (
        'attachment; filename="plano-44444444000144-20261008-formato-dataledger.txt"'
    )
    assert resposta["Content-Type"] == "text/plain; charset=utf-8"
    assert resposta.content == (
        "codigo;nome;codigo_pai;analitica;tipo;natureza\r\n"
        "1;Ativo;;N;ativo;devedora\r\n"
        "1.1;Caixa;1;S;ativo;devedora\r\n"
    ).encode("utf-8")


def test_exportacao_no_leiaute_da_ecd_nao_usa_ecd_nem_sped_no_nome(client, cenario, relogio):
    _entrar(client, "gestor-plano")
    _plano_de_duas_contas(cenario["empresa"])

    resposta = client.get(
        _url("plano_exportar", cenario["empresa"]),
        {"formato": "ecd", "data_alteracao": "2026-10-08"},
    )

    assert resposta.status_code == 200
    nome = re.search(r'filename="([^"]+)"', resposta["Content-Disposition"]).group(1)
    assert nome == "plano-44444444000144-20261008-registros-I050.txt"
    assert "ecd" not in nome.lower()
    assert "sped" not in nome.lower()
    assert resposta["Content-Type"] == "text/plain; charset=iso-8859-1"
    # Duas contas, uma linha I050 cada, com a data informada no DT_ALT, e nada além disso
    # (sem bloco J, sem termo).
    linhas = [linha for linha in resposta.content.decode("iso-8859-1").split("\r\n") if linha]
    assert len(linhas) == 2
    assert all(linha.startswith("|I050|08102026|") for linha in linhas)


def test_exportacao_ecd_sem_data_de_alteracao_e_recusada_com_mensagem(client, cenario):
    """O DataLedger não guarda a data da última alteração de cada conta e não inventa uma:
    a ECD exige a data informada pelo contador (DT_ALT)."""
    _entrar(client, "gestor-plano")
    _plano_de_duas_contas(cenario["empresa"])

    resposta = client.get(_url("plano_exportar", cenario["empresa"]), {"formato": "ecd"})

    assert resposta.status_code == 400
    assert "obrigatória para o I050" in resposta.content.decode()


def test_exportacao_no_leiaute_de_referencia_sai_com_cabecalho_0000(client, cenario, relogio):
    _entrar(client, "gestor-plano")
    _plano_de_duas_contas(cenario["empresa"])

    resposta = client.get(_url("plano_exportar", cenario["empresa"]), {"formato": "referencia"})

    assert resposta.status_code == 200
    assert "leiaute-com-separador.txt" in resposta["Content-Disposition"]
    # Forma canônica que o escritor grava: `|` no início e no fim de cada registro.
    assert resposta.content.decode("iso-8859-1").startswith(f"|0000|{CNPJ_DA_EMPRESA}|\r\n")


def test_exportacao_com_movimento_sem_periodo_e_recusada(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(
        _url("plano_exportar", cenario["empresa"]),
        {"formato": "proprio", "filtro": "com_movimento"},
    )

    assert resposta.status_code == 400
    assert "o filtro com movimento exige o período" in resposta.content.decode()


def test_exportacao_com_data_que_nao_existe_e_recusada(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(
        _url("plano_exportar", cenario["empresa"]),
        {
            "formato": "proprio",
            "filtro": "com_movimento",
            "inicio": "2026-02-31",
            "fim": "2026-03-01",
        },
    )

    assert resposta.status_code == 400


def test_exportacao_no_formato_excel_e_recusada_porque_so_importa(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(_url("plano_exportar", cenario["empresa"]), {"formato": "excel"})

    assert resposta.status_code == 400
    assert "não suportado para exportação" in resposta.content.decode()


def test_exportacao_com_parametro_nao_contratado_e_recusada(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(
        _url("plano_exportar", cenario["empresa"]), {"formato": "proprio", "xpto": "1"}
    )

    assert resposta.status_code == 400


def test_exportacao_sem_formato_escolhido_pede_a_escolha_em_vez_da_mensagem_tecnica(
    client, cenario
):
    _entrar(client, "gestor-plano")

    resposta = client.get(_url("plano_exportar", cenario["empresa"]), {"formato": ""})

    assert resposta.status_code == 400
    html = resposta.content.decode()
    assert "Escolha o formato do arquivo a exportar." in html
    assert "não suportado" not in html


def test_exportacao_registra_quem_exportou_com_sha_do_arquivo(client, cenario, relogio):
    _entrar(client, "gestor-plano")
    _plano_de_duas_contas(cenario["empresa"])

    resposta = client.get(_url("plano_exportar", cenario["empresa"]), {"formato": "proprio"})

    registro = RegistroAuditoria.objects.get(acao="plano_de_contas.exportado")
    assert registro.detalhes["sha256"] == hashlib.sha256(resposta.content).hexdigest()
    assert registro.detalhes["quantidade_contas"] == 2


# ---------------------------------------------------------------------------
# Modelo da planilha
# ---------------------------------------------------------------------------


def test_modelo_excel_baixa_um_xlsx_valido(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(_url("plano_modelo_excel", cenario["empresa"]))

    assert resposta.status_code == 200
    assert resposta["Content-Type"] == TIPO_XLSX
    assert resposta["Content-Disposition"] == 'attachment; filename="modelo-plano-de-contas.xlsx"'
    assert resposta.content[:2] == b"PK"
    pasta = load_workbook(io.BytesIO(resposta.content))
    assert "plano" in pasta.sheetnames
    assert "instrucoes" in pasta.sheetnames
    cabecalho = [celula.value for celula in pasta["plano"][1]][: len(CABECALHO)]
    assert cabecalho == list(CABECALHO)


def test_modelo_excel_nao_aceita_parametro_na_url(client, cenario):
    _entrar(client, "gestor-plano")

    resposta = client.get(_url("plano_modelo_excel", cenario["empresa"]), {"x": "1"})

    assert resposta.status_code == 400


# ---------------------------------------------------------------------------
# Permissões, isolamento, anônimo, livro-caixa, CSRF
# ---------------------------------------------------------------------------

ROTAS_DE_ESCRITA = [
    ("plano_importar", "POST"),
    ("plano_importar_aplicar", "POST"),
]
ROTAS = [
    ("plano_importar", "GET"),
    ("plano_importar", "POST"),
    ("plano_importar_aplicar", "POST"),
    ("plano_modelo_excel", "GET"),
    ("plano_exportar", "GET"),
]


def _chamar(client, empresa, nome, metodo):
    url = _url(nome, empresa)
    if metodo == "GET":
        return client.get(url)
    return client.post(url, {"formato": "proprio", "arquivo": _arquivo(PROPRIO, "p.txt")})


@pytest.mark.parametrize("nome,metodo", ROTAS)
def test_cliente_nao_passa_em_nenhuma_rota_do_plano_em_arquivo(client, cenario, nome, metodo):
    _entrar(client, "cliente-plano")

    assert _chamar(client, cenario["empresa"], nome, metodo).status_code == 403


@pytest.mark.parametrize("nome,metodo", ROTAS_DE_ESCRITA)
def test_papel_sem_edicao_do_plano_nao_importa_nem_aplica(client, cenario, nome, metodo):
    _entrar(client, "paralegal-plano")

    resposta = _chamar(client, cenario["empresa"], nome, metodo)

    assert resposta.status_code == 403
    assert _contas(cenario["empresa"]) == set()


def test_paralegal_nao_baixa_modelo_mas_exporta_o_plano(client, cenario):
    _entrar(client, "paralegal-plano")

    assert client.get(_url("plano_modelo_excel", cenario["empresa"])).status_code == 403
    assert client.get(_url("plano_exportar", cenario["empresa"])).status_code == 200


@pytest.mark.parametrize("usuario", ["gestor-plano", "analista-plano", "financeiro-plano"])
def test_quem_escreve_o_plano_acessa_as_quatro_telas(client, cenario, usuario):
    _entrar(client, usuario)
    empresa = cenario["empresa"]

    assert client.get(_url("plano_importar", empresa)).status_code == 200
    assert client.get(_url("plano_exportar", empresa)).status_code == 200
    assert client.get(_url("plano_modelo_excel", empresa)).status_code == 200


@pytest.mark.parametrize("nome,metodo", ROTAS)
def test_empresa_de_outro_escritorio_responde_404_nos_dois_sentidos(client, cenario, nome, metodo):
    """Isolamento entre escritórios, nos dois sentidos: 404, nunca 403 (não confirma existência)."""
    _entrar(client, "gestor-plano")
    assert _chamar(client, cenario["de_outro_escritorio"], nome, metodo).status_code == 404
    client.logout()
    _entrar(client, "gestor-outro-escritorio")
    assert _chamar(client, cenario["empresa"], nome, metodo).status_code == 404


@pytest.mark.parametrize("nome,metodo", ROTAS)
def test_anonimo_vai_para_o_login(client, cenario, nome, metodo):
    resposta = _chamar(client, cenario["empresa"], nome, metodo)

    assert resposta.status_code == 302
    assert "login" in resposta["Location"]


@pytest.mark.parametrize("nome,metodo", ROTAS)
def test_empresa_em_livro_caixa_e_recusada_em_toda_rota(
    client, cenario, empresa_livro_caixa, nome, metodo
):
    _entrar(client, "gestor-plano")

    resposta = _chamar(client, empresa_livro_caixa, nome, metodo)

    assert resposta.status_code == 403
    assert MENSAGEM_RECUSA_CONTABILIDADE_LIVRO_CAIXA in resposta.content.decode()
    assert _contas(empresa_livro_caixa) == set()


def test_post_sem_token_csrf_e_recusado(cenario):
    cliente = Client(enforce_csrf_checks=True)
    _entrar(cliente, "gestor-plano")

    resposta = cliente.post(
        _url("plano_importar", cenario["empresa"]),
        {"arquivo": _arquivo(PROPRIO, "p.txt"), "formato": "proprio", "politica": "so_acrescentar"},
    )

    # 403 por CSRF, não por permissão: o gestor tem o papel, então a recusa é do token.
    assert resposta.status_code == 403
    assert "CSRF" in resposta.content.decode()
    assert _contas(cenario["empresa"]) == set()


def test_telas_do_plano_em_arquivo_passam_as_guardas_de_acessibilidade(client, cenario):
    _entrar(client, "gestor-plano")
    empresa = cenario["empresa"]
    _plano_de_duas_contas(empresa)

    importar = client.get(_url("plano_importar", empresa)).content.decode()
    previa = _previa(client, empresa, PROPRIO, "plano.txt", "proprio").content.decode()
    exportar = client.get(_url("plano_exportar", empresa)).content.decode()

    for html in (importar, previa, exportar):
        assert_pagina_acessivel(html)


def test_plano_de_contas_mostra_importar_so_a_quem_escreve(client, cenario):
    _entrar(client, "gestor-plano")
    escrita = client.get(_url("plano_de_contas", cenario["empresa"])).content.decode()
    client.logout()
    _entrar(client, "paralegal-plano")
    leitura = client.get(_url("plano_de_contas", cenario["empresa"])).content.decode()

    assert "Importar plano" in escrita
    assert "Importar plano" not in leitura
    assert "Exportar plano" in escrita
    assert "Exportar plano" in leitura


# ---------------------------------------------------------------------------
# Inventários: a tela enxerga o mesmo que o núcleo
# ---------------------------------------------------------------------------


def test_rotulos_da_tela_cobrem_exatamente_os_formatos_e_as_politicas_do_nucleo():
    assert set(FORMATOS_DE_IMPORTACAO) == set(LEITORES)
    assert set(FORMATOS_DE_EXPORTACAO) == set(ESCRITORES)
    assert set(POLITICAS_DE_IMPORTACAO_NA_TELA) == set(POLITICAS)
    assert set(FILTROS_DE_EXPORTACAO_NA_TELA) == set(FILTROS_DE_EXPORTACAO)
