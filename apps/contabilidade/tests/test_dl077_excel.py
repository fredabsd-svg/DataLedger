"""DL-077, fatia 1, frente B: importação do plano por planilha `.xlsx`
(`formatos/excel.py`). Critério 6 da DL-077, inteiro: macro, fórmula sem valor,
código numérico, arquivo acima do limite, aba fora do modelo e `.xls`/ZIP malicioso.

As planilhas são geradas AQUI com openpyxl, a partir de dados sintéticos. Os pacotes
hostis são montados byte a byte, a partir de um .xlsx válido. Nada vem de cliente real.
"""

import io
import zipfile
from datetime import datetime

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.auditoria.models import RegistroAuditoria
from apps.contabilidade.intercambio import leitura
from apps.contabilidade.intercambio.canonico import NIVEL_AVISO, NIVEL_ERRO, IntercambioRecusado
from apps.contabilidade.intercambio.formatos import ESCRITORES, LEITORES, excel
from apps.contabilidade.intercambio.formatos.excel import gerar_modelo, ler
from apps.contabilidade.intercambio.formatos.proprio import CABECALHO
from apps.contabilidade.intercambio.leitura import (
    TAMANHO_MAXIMO_ARQUIVO_BYTES,
    ArquivoGrandeDemais,
    ler_arquivo,
)
from apps.contabilidade.models import Conta
from apps.empresas.models import Empresa
from apps.tenancy.models import Escritorio, Papel, VinculoUsuarioEscritorio

pytestmark = pytest.mark.django_db

SENHA = "senha-forte-123"
CAB = list(CABECALHO)
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAIN_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
MAIN_XLSM = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
MAIN_XLTX = "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml"


def _pasta(linhas, *, abas_extras=(), aba="plano"):
    """Gera um .xlsx com as linhas dadas na aba `aba`. Linha `[]` vira linha em branco."""
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = aba
    for linha in linhas:
        planilha.append(linha)
    for nome in abas_extras:
        pasta.create_sheet(nome)
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def _reescrever_pacote(conteudo, *, trocar=None, acrescentar=None):
    """Copia um .xlsx parte a parte, trocando e acrescentando partes."""
    origem = zipfile.ZipFile(io.BytesIO(conteudo))
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as destino:
        for parte in origem.infolist():
            dados = origem.read(parte.filename)
            if trocar and parte.filename in trocar:
                dados = trocar[parte.filename](dados)
            destino.writestr(parte.filename, dados)
        for nome, dados in (acrescentar or {}).items():
            destino.writestr(nome, dados)
    return saida.getvalue()


def _tipo_de_pasta(de, para):
    """Troca o tipo de conteúdo de `xl/workbook.xml` em [Content_Types].xml."""
    return lambda dados: dados.replace(de.encode(), para.encode())


def _erros(resultado):
    return [o for o in resultado.ocorrencias if o.nivel == NIVEL_ERRO]


# -----------------------------------------------------------------------------
# Leitura da planilha
# -----------------------------------------------------------------------------


def test_planilha_valida_lida_com_numero_de_linha_da_planilha_e_linha_em_branco_ignorada():
    resultado = ler(
        _pasta(
            [
                CAB,
                ["1", "Ativo", "", "N", "ativo", "devedora"],
                [],  # linha 3 em branco: ignorada
                ["1.1", "Caixa", "1", "S", "", ""],
            ]
        )
    )

    assert not resultado.tem_erro, resultado.ocorrencias
    assert [(c.codigo, c.linha) for c in resultado.contas] == [("1", 2), ("1.1", 4)]
    assert resultado.contas[1].codigo_pai == "1"
    assert resultado.contas[1].analitica is True
    assert resultado.formato == "excel"


def test_aba_extra_e_ignorada_com_aviso_e_nao_tem_erro():
    resultado = ler(
        _pasta(
            [CAB, ["1", "Ativo", "", "N", "ativo", "devedora"]],
            abas_extras=["instrucoes"],
        )
    )

    assert not resultado.tem_erro
    avisos = [o for o in resultado.ocorrencias if o.nivel == NIVEL_AVISO]
    assert any(o.campo == "aba" and "instrucoes" in o.mensagem for o in avisos)
    assert len(resultado.contas) == 1


def test_aba_plano_ausente_e_erro_e_nada_e_lido():
    resultado = ler(_pasta([CAB, ["1", "Ativo", "", "N", "ativo", "devedora"]], aba="dados"))

    assert resultado.contas == []
    assert any(o.campo == "aba" and o.nivel == NIVEL_ERRO for o in resultado.ocorrencias)


def test_cabecalho_fora_do_modelo_e_erro_e_nenhuma_linha_e_lida():
    trocado = ["codigo", "nome", "analitica", "codigo_pai", "tipo", "natureza"]

    resultado = ler(_pasta([trocado, ["1", "Ativo", "N", "", "ativo", "devedora"]]))

    assert resultado.contas == []
    assert [(o.linha, o.campo) for o in _erros(resultado)] == [(1, "cabecalho")]


def test_coluna_alem_da_sexta_com_conteudo_e_erro():
    resultado = ler(_pasta([CAB, ["1", "Ativo", "", "N", "ativo", "devedora", "sobra"]]))

    assert resultado.contas == []
    assert any(
        o.linha == 2 and o.campo == "coluna 7" and o.nivel == NIVEL_ERRO
        for o in resultado.ocorrencias
    )


def test_formula_sem_valor_calculado_e_erro_e_a_linha_nao_entra():
    resultado = ler(
        _pasta(
            [
                CAB,
                ["1", '=CONCAT("Ativo","")', "", "N", "ativo", "devedora"],
            ]
        )
    )

    assert resultado.contas == []
    erro = _erros(resultado)[0]
    assert (erro.linha, erro.campo) == (2, "nome")
    assert "fórmula sem valor calculado" in erro.mensagem


def test_celula_com_erro_de_formula_salvo_e_recusada():
    """Erro salvo no arquivo (ex.: #DIV/0!) não é um valor: a coluna vira erro nomeado."""
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = "plano"
    planilha.append(CAB)
    planilha.append(["1", "Ativo", "", "N", "ativo", "devedora"])
    planilha.append(["1.1", "Caixa", "1", "S", "x", ""])
    celula = planilha["E3"]
    celula.value = "#DIV/0!"
    celula.data_type = "e"
    saida = io.BytesIO()
    pasta.save(saida)

    resultado = ler(saida.getvalue())

    assert [(c.codigo) for c in resultado.contas] == ["1"]
    assert any(
        o.linha == 3 and o.campo == "tipo" and "erro de fórmula" in o.mensagem
        for o in _erros(resultado)
    )


def test_codigo_numerico_e_recusado_e_nao_e_convertido():
    """O Excel grava `1.1` como 1,1. O importador não converte: recusa a célula. Se
    convertesse, o código sairia `1.1` ou `1.10` sem ninguém perceber."""
    resultado = ler(_pasta([CAB, [1.1, "Caixa", "1", "S", "", ""]]))

    assert resultado.contas == []
    erro = _erros(resultado)[0]
    assert (erro.linha, erro.campo) == (2, "codigo")
    assert "Formate a coluna" in erro.mensagem and "Texto" in erro.mensagem


def test_codigo_pai_numerico_tambem_e_recusado():
    resultado = ler(_pasta([CAB, ["1.1", "Caixa", 1.1, "S", "", ""]]))

    assert resultado.contas == []
    assert any(o.campo == "codigo_pai" and o.nivel == NIVEL_ERRO for o in resultado.ocorrencias)


def test_data_na_coluna_de_codigo_e_recusada():
    resultado = ler(_pasta([CAB, [datetime(2026, 1, 31), "Caixa", "", "S", "", ""]]))

    assert resultado.contas == []
    erro = _erros(resultado)[0]
    assert erro.campo == "codigo" and "como data" in erro.mensagem


def test_numero_no_nome_e_recusado_em_vez_de_convertido():
    resultado = ler(_pasta([CAB, ["1", 2026, "", "N", "ativo", "devedora"]]))

    assert resultado.contas == []
    assert any(o.campo == "nome" and o.nivel == NIVEL_ERRO for o in resultado.ocorrencias)


def test_codigo_repetido_na_planilha_e_erro_pelas_mesmas_regras_do_formato_proprio():
    resultado = ler(
        _pasta(
            [
                CAB,
                ["1", "Ativo", "", "N", "ativo", "devedora"],
                ["1", "Ativo de novo", "", "N", "ativo", "devedora"],
            ]
        )
    )

    assert [c.codigo for c in resultado.contas] == ["1"]
    assert any(o.linha == 3 and o.campo == "codigo" for o in _erros(resultado))


def test_planilha_so_com_cabecalho_nao_tem_erro_e_nao_tem_conta():
    resultado = ler(_pasta([CAB]))

    assert not resultado.tem_erro
    assert resultado.contas == []


# -----------------------------------------------------------------------------
# Recusas do pacote (`.xls`, CSV renomeado, macro, bomba, pacote danificado)
# -----------------------------------------------------------------------------


def test_xls_formato_binario_antigo_e_recusado():
    xls = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512

    with pytest.raises(IntercambioRecusado, match=r"\.xls"):
        ler(xls)


def test_csv_renomeado_para_xlsx_e_recusado():
    csv = "codigo;nome;codigo_pai;analitica;tipo;natureza\r\n1;Ativo;;N;ativo;devedora\r\n"

    with pytest.raises(IntercambioRecusado, match="não é uma planilha .xlsx"):
        ler(csv.encode("utf-8"))


def test_zip_que_nao_e_pasta_de_trabalho_e_recusado():
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as pacote:
        pacote.writestr("dados.txt", "nada")

    with pytest.raises(IntercambioRecusado, match="Content_Types"):
        ler(saida.getvalue())


def test_xlsm_pelo_tipo_de_conteudo_e_recusado_mesmo_sem_parte_vba():
    """Mutação (c): aceitar `.xlsm` derruba este teste. Sem a parte VBA, a única marca
    de macro é o tipo declarado da pasta."""
    xlsm = _reescrever_pacote(
        _pasta([CAB]),
        trocar={"[Content_Types].xml": _tipo_de_pasta(MAIN_XLSX, MAIN_XLSM)},
    )

    with pytest.raises(IntercambioRecusado, match="xlsm"):
        ler(xlsm)


def test_parte_vba_e_recusada_mesmo_com_extensao_xlsx():
    com_macro = _reescrever_pacote(
        _pasta([CAB, ["1", "Ativo", "", "N", "ativo", "devedora"]]),
        acrescentar={"xl/vbaProject.bin": b"\x00" * 64},
    )

    with pytest.raises(IntercambioRecusado, match="macros"):
        ler(com_macro)


def test_modelo_de_pasta_xltx_e_recusado_por_nao_ser_pasta_de_trabalho_comum():
    xltx = _reescrever_pacote(
        _pasta([CAB]),
        trocar={"[Content_Types].xml": _tipo_de_pasta(MAIN_XLSX, MAIN_XLTX)},
    )

    with pytest.raises(IntercambioRecusado, match="pasta de trabalho .xlsx comum"):
        ler(xltx)


def test_pacote_danificado_e_recusado():
    with pytest.raises(IntercambioRecusado, match="danificado"):
        ler(b"PK\x03\x04" + b"lixo que nao e um zip valido" * 10)


def test_bomba_de_descompactacao_e_recusada_pelo_tamanho_medido_no_conteudo(monkeypatch):
    """O limite vale para o DESCOMPACTADO, medido lendo as partes. Um pacote de poucos KB
    que se expande além do limite é recusado antes de o openpyxl abri-lo. O teste baixa
    o limite para 1 MB, para não criar 200 MB de teste."""
    monkeypatch.setattr(excel, "TAMANHO_MAXIMO_DESCOMPACTADO_BYTES", 1024 * 1024)
    bomba = _reescrever_pacote(_pasta([CAB]), acrescentar={"xl/extra.bin": b"\x00" * (2 << 20)})

    assert len(bomba) < 100 * 1024  # compactado, é pequeno: a expansão é o problema
    with pytest.raises(ArquivoGrandeDemais, match="descompactada"):
        ler(bomba)


def test_pacote_acima_do_limite_de_partes_e_recusado(monkeypatch):
    monkeypatch.setattr(excel, "MAXIMO_DE_ENTRADAS_NO_PACOTE", 3)
    muitas = _reescrever_pacote(
        _pasta([CAB]),
        acrescentar={f"xl/extra{i}.bin": b"x" for i in range(5)},
    )

    with pytest.raises(IntercambioRecusado, match="partes internas"):
        ler(muitas)


def test_arquivo_acima_do_limite_de_bytes_e_recusado_antes_de_abrir():
    with pytest.raises(ArquivoGrandeDemais):
        ler_arquivo("excel", b"PK\x03\x04" + b"0" * TAMANHO_MAXIMO_ARQUIVO_BYTES)


def test_planilha_acima_do_limite_de_linhas_e_recusada(monkeypatch):
    monkeypatch.setattr(leitura, "MAXIMO_DE_LINHAS", 2)
    planilha = _pasta([CAB] + [[str(i), "x", "", "N", "ativo", "devedora"] for i in range(5)])

    with pytest.raises(ArquivoGrandeDemais, match="linhas"):
        ler(planilha)


def test_leitura_pelo_ponto_de_entrada_comum_marca_formato_e_sha256():
    conteudo = _pasta([CAB, ["1", "Ativo", "", "N", "ativo", "devedora"]])

    resultado = ler_arquivo("excel", conteudo, nome_arquivo="plano.xlsx")

    assert resultado.formato == "excel"
    assert len(resultado.sha256) == 64
    assert resultado.nome_arquivo == "plano.xlsx"


# -----------------------------------------------------------------------------
# Registro, modelo para baixar
# -----------------------------------------------------------------------------


def test_excel_so_e_leitor_nao_tem_escritor():
    assert LEITORES["excel"] is ler
    assert "excel" not in ESCRITORES


def test_modelo_tem_aba_plano_cabecalho_coluna_de_codigo_como_texto_e_instrucoes():
    modelo = load_workbook(io.BytesIO(gerar_modelo()))

    assert modelo.sheetnames == ["plano", "instrucoes"]
    plano = modelo["plano"]
    assert [plano.cell(row=1, column=c).value for c in range(1, 7)] == CAB
    assert plano["A2"].number_format == "@"  # codigo
    assert plano["C2"].number_format == "@"  # codigo_pai
    assert plano["B2"].number_format != "@"
    textos = [c.value for c in modelo["instrucoes"]["A"] if c.value]
    assert "codigo" in textos and "natureza" in textos


def test_modelo_vazio_le_sem_erro_e_sem_conta():
    resultado = ler(gerar_modelo())

    assert not resultado.tem_erro
    assert resultado.contas == []


def test_codigo_digitado_no_modelo_como_texto_e_lido_sem_conversao():
    modelo = load_workbook(io.BytesIO(gerar_modelo()))
    plano = modelo["plano"]
    # A aba já tem o cabeçalho; a linha 2 é a que o contador digita.
    for coluna, valor in enumerate(["1.1", "Caixa", "", "S", "", ""], start=1):
        plano.cell(row=2, column=coluna, value=valor)
    saida = io.BytesIO()
    modelo.save(saida)

    resultado = ler(saida.getvalue())

    assert not resultado.tem_erro, resultado.ocorrencias
    assert [c.codigo for c in resultado.contas] == ["1.1"]


# -----------------------------------------------------------------------------
# API: prévia, aplicação, modelo, isolamento e permissão
# -----------------------------------------------------------------------------

PLANO_VALIDO = _pasta(
    [
        CAB,
        ["1", "Ativo", "", "N", "ativo", "devedora"],
        ["1.1", "Caixa", "1", "S", "", ""],
    ]
)
PLANO_COM_MACRO = _reescrever_pacote(PLANO_VALIDO, acrescentar={"xl/vbaProject.bin": b"\x00"})


def _usuario(username, escritorio, papel):
    usuario = get_user_model().objects.create_user(
        username=username, email=f"{username}@escritorio.com.br", password=SENHA
    )
    VinculoUsuarioEscritorio.objects.create(usuario=usuario, escritorio=escritorio, papel=papel)
    return usuario


@pytest.fixture
def cenario(client):
    escritorio_a = Escritorio.objects.create(nome="Escritório Excel A", cnpj="55555555000155")
    escritorio_b = Escritorio.objects.create(nome="Escritório Excel B", cnpj="66666666000166")
    empresa = Empresa.objects.create(
        escritorio=escritorio_a, razao_social="Empresa Excel Ltda", cnpj="77777777000177"
    )
    outra = Empresa.objects.create(
        escritorio=escritorio_b, razao_social="Outra Excel Ltda", cnpj="88888888000188"
    )
    _usuario("analista-excel", escritorio_a, Papel.ANALISTA)
    _usuario("paralegal-excel", escritorio_a, Papel.PARALEGAL)
    _usuario("cliente-excel", escritorio_a, Papel.CLIENTE)
    _usuario("gestor-b-excel", escritorio_b, Papel.GESTOR)
    return {"empresa": empresa, "outra": outra}


def _entrar(client, username):
    assert client.login(username=username, password=SENHA)


def _previa_url(empresa_id):
    return reverse("contabilidade:plano-importacao-previa", args=[empresa_id])


def _aplicar_url(empresa_id):
    return reverse("contabilidade:plano-importacao-aplicar", args=[empresa_id])


def _modelo_url(empresa_id):
    return reverse("contabilidade:plano-modelo-excel", args=[empresa_id])


def _planilha(conteudo, nome="plano.xlsx"):
    return SimpleUploadedFile(nome, conteudo, content_type=TIPO_XLSX)


def _previa(client, empresa, conteudo=PLANO_VALIDO, nome="plano.xlsx"):
    return client.post(
        _previa_url(empresa.id),
        {"arquivo": _planilha(conteudo, nome), "formato": "excel"},
    )


def test_api_previa_de_planilha_xlsx_valida_aceita(client, cenario):
    _entrar(client, "analista-excel")

    resposta = _previa(client, cenario["empresa"])

    assert resposta.status_code == 200, resposta.content
    corpo = resposta.json()
    assert corpo["formato"] == "excel"
    assert corpo["pode_aplicar"] is True
    assert corpo["contagens"]["criar"] == 2
    assert Conta.objects.filter(empresa=cenario["empresa"]).count() == 0


@pytest.mark.parametrize("nome", ["plano.xlsm", "plano.xls", "plano.csv"])
def test_api_previa_recusa_nome_que_nao_e_xlsx_com_400(client, cenario, nome):
    _entrar(client, "analista-excel")

    resposta = _previa(client, cenario["empresa"], nome=nome)

    assert resposta.status_code == 400
    assert "arquivo" in resposta.json()


def test_api_previa_recusa_macro_com_nome_xlsx_com_400_e_nada_grava(client, cenario):
    _entrar(client, "analista-excel")

    resposta = _previa(client, cenario["empresa"], conteudo=PLANO_COM_MACRO)

    assert resposta.status_code == 400
    assert "macros" in resposta.json()["arquivo"][0]
    assert Conta.objects.filter(empresa=cenario["empresa"]).count() == 0


def test_api_previa_excel_acima_do_limite_responde_413(client, cenario):
    _entrar(client, "analista-excel")
    grande = b"PK\x03\x04" + b"0" * TAMANHO_MAXIMO_ARQUIVO_BYTES

    resposta = _previa(client, cenario["empresa"], conteudo=grande)

    assert resposta.status_code == 413


def test_api_aplicacao_de_planilha_grava_com_o_token_da_previa_e_registra_trilha(client, cenario):
    empresa = cenario["empresa"]
    _entrar(client, "analista-excel")
    previa = _previa(client, empresa).json()

    resposta = client.post(
        _aplicar_url(empresa.id),
        {
            "arquivo": _planilha(PLANO_VALIDO),
            "formato": "excel",
            "sha256": previa["sha256"],
            "assinatura": previa["assinatura"],
        },
    )

    assert resposta.status_code == 200, resposta.content
    assert resposta.json()["criadas"] == 2
    assert Conta.objects.filter(empresa=empresa).count() == 2
    trilha = RegistroAuditoria.objects.get(acao="plano_de_contas.importado")
    assert trilha.detalhes["formato"] == "excel"
    assert trilha.detalhes["sha256"] == previa["sha256"]


@pytest.mark.parametrize("usuario", ["paralegal-excel", "cliente-excel"])
def test_api_previa_de_planilha_recusa_quem_nao_escreve_o_plano(client, cenario, usuario):
    _entrar(client, usuario)

    assert _previa(client, cenario["empresa"]).status_code == 403


def test_api_modelo_baixa_xlsx_como_anexo_e_o_arquivo_abre(client, cenario):
    _entrar(client, "analista-excel")

    resposta = client.get(_modelo_url(cenario["empresa"].id))

    assert resposta.status_code == 200, resposta.content
    assert resposta.headers["Content-Type"] == TIPO_XLSX
    assert resposta.headers["Content-Disposition"] == (
        'attachment; filename="modelo-plano-de-contas.xlsx"'
    )
    modelo = load_workbook(io.BytesIO(resposta.content))
    assert modelo.sheetnames == ["plano", "instrucoes"]


@pytest.mark.parametrize("usuario", ["paralegal-excel", "cliente-excel"])
def test_api_modelo_recusa_quem_nao_escreve_o_plano(client, cenario, usuario):
    _entrar(client, usuario)

    assert client.get(_modelo_url(cenario["empresa"].id)).status_code == 403


def test_api_modelo_de_empresa_de_outro_escritorio_responde_404(client, cenario):
    _entrar(client, "analista-excel")

    assert client.get(_modelo_url(cenario["outra"].id)).status_code == 404


def test_api_modelo_recusa_parametro_que_a_rota_nao_le(client, cenario):
    _entrar(client, "analista-excel")

    assert client.get(_modelo_url(cenario["empresa"].id), {"xpto": "1"}).status_code == 400
